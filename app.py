"""Librería MM — Sistema de Inventario, Ventas/Caja, Deudores y Reportes.
Stack: Streamlit + Pandas + Excel local (data/*.xlsx).
Moneda: Guaraníes (Gs.) sin decimales. Precio_Sugerido = Costo * 1.5.
Ejecución: streamlit run app.py
"""
from __future__ import annotations

import hashlib
import html
import io
import os
import re
import threading
import time
from datetime import date, datetime
from pathlib import Path

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import drive_sync  # almacenamiento persistente (Drive en nube, local en PC)

# ---------------------------------------------------------------- Config
APP_TITULO = "📚 Librería MM — Gestión"
BASE = Path(__file__).resolve().parent
DATA_DIR = BASE / "data"
INV_PATH = DATA_DIR / "inventario.xlsx"
VENTAS_PATH = DATA_DIR / "ventas.xlsx"
DEUD_PATH = DATA_DIR / "deudores.xlsx"
USU_PATH = DATA_DIR / "usuarios.xlsx"
CAT_PATH = DATA_DIR / "categorias.txt"
CAT_HIDE_PATH = DATA_DIR / "categorias_ocultas.txt"
CATEGORIA_DEFECTO = "Otros"  # Categoría de respaldo: nunca se puede borrar

MARGEN = 1.5  # CORRECCIÓN 4: Precio_Sugerido = Costo de Compra * 1.5 (nunca desde venta)
CATEGORIAS_DEFAULT = ["Libros", "Útiles", "Fotocopias", "Servicios", "Juguetes", "Regalería", "Otros"]
METODOS_PAGO = ["Efectivo", "Tarjeta", "Transferencia", "Otro"]
CAJAS = ["Caja 1", "Caja 2"]

COL_INV = ["SKU", "Nombre", "Categoria", "Costo", "Precio_Sugerido",
           "Precio_Actual", "Stock_Actual", "Stock_Minimo", "Fecha_Actualizacion"]
COL_VENTAS = ["ID_Venta", "Fecha", "SKU", "Producto", "Categoria", "Cantidad",
              "Precio_Unitario", "Total", "Metodo_Pago", "Caja", "Notas",
              "Vendedor_Doc", "Vendedor_Nombre"]
COL_DEUD = ["ID", "Cliente", "Telefono", "Concepto", "Monto_Total",
            "Monto_Pagado", "Saldo", "Fecha_Prestamo", "Fecha_Vencimiento", "Estado"]
COL_USU = ["Documento", "Nombre", "Password_Hash", "Rol", "Activo", "Fecha_Creacion"]

ROLES = ["SuperUsuario", "Admin", "Cajero"]
# Usuario SuperUsuario inicial (se crea solo si no existe)
SUPER_DOC = "5729679"
SUPER_NOMBRE = "Cesar Verza"
SUPER_PASS = "Nicolas08.L"

FILL_ENCABEZADO = PatternFill("solid", fgColor="1F4E79")
FONT_ENCABEZADO = Font(bold=True, color="FFFFFF", size=11)

st.set_page_config(page_title="Librería MM", page_icon="📚", layout="wide")

# Las descargas nativas de Streamlit ("Download as CSV" de tablas/gráficos) usan
# coma como separador y Excel en español las abre en una sola columna. Se ocultan
# y en su lugar cada sección ofrece botones propios: .xlsx o CSV con `;` y BOM.
# Además: animación de entrada para gráficos/métricas y superficies neutras que
# mantienen contraste en modo claro Y oscuro (sin deslumbrar).
st.markdown(
    "<style>"
    "div[data-testid='stElementToolbar']{display:none;}"
    "@keyframes mm-in{from{opacity:0;transform:translateY(10px);}to{opacity:1;transform:none;}}"
    "[data-testid='stVegaLiteChart'],[data-testid='stMetric']{animation:mm-in .45s ease both;}"
    "[data-testid='stMetric']{background:rgba(127,127,127,.09);"
    "border:1px solid rgba(127,127,127,.22);border-radius:14px;padding:12px 14px;}"
    "[data-testid='stMetric'] label{font-weight:600;}"
    "[data-testid='stMetric'] [data-testid='stMetricValue']{font-variant-numeric:tabular-nums;}"
    "div[data-testid='stAlert']{border-radius:12px;}"
    ".stTabs [data-baseweb='tab']{border-radius:10px 10px 0 0;}"
    "button[kind='primary']{border-radius:10px;}"
    "@media (prefers-color-scheme:dark){"
    "[data-testid='stMetric']{background:rgba(255,255,255,.05);"
    "border-color:rgba(255,255,255,.14);box-shadow:none;}"
    "div[data-testid='stAlert']{filter:saturate(.92) brightness(.96);}"
    "}</style>",
    unsafe_allow_html=True)


def csv_puntoycoma_bytes(df: pd.DataFrame) -> bytes:
    """CSV con `;` como separador y BOM UTF-8 para que Excel lo abra en columnas."""
    return df.to_csv(index=False, sep=";").encode("utf-8-sig")

# --------------------------------------------- Utilidades: formato Gs. entero
def a_entero(x) -> int:
    """Convierte a entero sin decimales (1000.00 -> 1000)."""
    try:
        return int(round(float(x)))
    except (TypeError, ValueError):
        return 0


def fmt_gs(x) -> str:
    """Formato moneda paraguaya: Gs. 1.000.000 (punto = miles, sin decimales)."""
    try:
        n = int(round(float(x)))
    except (TypeError, ValueError):
        n = 0
    return "Gs. " + f"{n:,}".replace(",", ".")


def ahora_str() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M")


def recalcular_sugerido(costo) -> int:
    """CORRECCIÓN 4 — Fórmula estricta: Precio Sugerido = Costo de Compra * 1.5.

    Siempre parte del Costo (compra al proveedor). Nunca usa Precio_Actual.
    Devuelve entero (Gs. no usa céntimos).
    """
    try:
        return int(round(float(costo) * MARGEN))
    except (TypeError, ValueError):
        return 0


# ------------------------------------------------- Categorías dinámicas (2)
@st.cache_data(ttl=300)
def _categorias_disco() -> tuple[list[str], list[str]]:
    guardadas = []
    if CAT_PATH.exists():
        try:
            guardadas = [l.strip() for l in CAT_PATH.read_text(encoding="utf-8").splitlines() if l.strip()]
        except OSError:
            guardadas = []
    ocultas = []
    if CAT_HIDE_PATH.exists():
        try:
            ocultas = [l.strip() for l in CAT_HIDE_PATH.read_text(encoding="utf-8").splitlines() if l.strip()]
        except OSError:
            ocultas = []
    return guardadas, ocultas


def _categorias_guardadas() -> list[str]:
    return list(_categorias_disco()[0])


def _categorias_ocultas() -> list[str]:
    """Categorías predeterminadas que el usuario borró (para no volver a mostrarlas)."""
    return list(_categorias_disco()[1])


def guardar_nueva_categoria(nombre: str) -> bool:
    nombre = (nombre or "").strip()
    if not nombre:
        return False
    actuales = obtener_categorias(pd.DataFrame())
    if nombre.lower() in {c.lower() for c in actuales}:
        return False
    DATA_DIR.mkdir(exist_ok=True)
    # Si estaba oculta por un borrado previo, rehabilitarla
    ocultas = [c for c in _categorias_ocultas() if c.lower() != nombre.lower()]
    CAT_HIDE_PATH.write_text("\n".join(sorted(set(ocultas))), encoding="utf-8")
    extras = _categorias_guardadas()
    extras.append(nombre)
    CAT_PATH.write_text("\n".join(sorted(set(extras))), encoding="utf-8")
    _categorias_disco.clear()
    return True


def eliminar_categoria(nombre: str) -> tuple[bool, int]:
    """Borra una categoría con borrado seguro.

    Regla de integridad: los productos que la tenían se reasignan a "Otros".
    Devuelve (ok, cantidad_reasignada). "Otros" no se puede borrar.
    """
    nombre = (nombre or "").strip()
    if not nombre or nombre.lower() == CATEGORIA_DEFECTO.lower():
        return False, 0
    # 1. Reasignar productos afectados a la categoría por defecto
    n = 0
    df = cargar_inventario()
    if len(df):
        mask = df["Categoria"].astype(str).str.strip().str.lower() == nombre.lower()
        n = int(mask.sum())
        if n:
            df.loc[mask, "Categoria"] = CATEGORIA_DEFECTO
            guardar_inventario(df)
    # 2. Quitar de personalizadas y ocultar si era predeterminada
    DATA_DIR.mkdir(exist_ok=True)
    extras = [c for c in _categorias_guardadas() if c.lower() != nombre.lower()]
    CAT_PATH.write_text("\n".join(sorted(set(extras))), encoding="utf-8")
    if nombre.lower() in {c.lower() for c in CATEGORIAS_DEFAULT}:
        ocultas = _categorias_ocultas() + [nombre]
        CAT_HIDE_PATH.write_text("\n".join(sorted(set(ocultas))), encoding="utf-8")
    _categorias_disco.clear()
    cargar_inventario.clear()
    return True, n


def obtener_categorias(inv: pd.DataFrame) -> list[str]:
    """Lista unificada: default + guardadas en disco + las que ya existen en el inventario.

    Excluye las borradas por el usuario; "Otros" siempre está presente.
    """
    ocultas = {c.lower() for c in _categorias_ocultas()}
    cats: list[str] = [c for c in list(CATEGORIAS_DEFAULT) + _categorias_guardadas()
                       if c.lower() not in ocultas]
    if inv is not None and len(inv) and "Categoria" in inv.columns:
        cats += [str(c).strip() for c in inv["Categoria"].dropna().unique().tolist() if str(c).strip()]
    vistas, vistos = [], set()
    for c in cats:
        if c.lower() not in vistos:
            vistas.append(c)
            vistos.add(c.lower())
    if CATEGORIA_DEFECTO.lower() not in vistos:
        vistas.append(CATEGORIA_DEFECTO)
    return sorted(vistas)


def df_vacio_inv() -> pd.DataFrame:
    return pd.DataFrame(columns=COL_INV)


def df_vacio_ventas() -> pd.DataFrame:
    return pd.DataFrame(columns=COL_VENTAS)


def df_vacio_deud() -> pd.DataFrame:
    return pd.DataFrame(columns=COL_DEUD)


def _leer_excel_seguro(path: Path, hoja: str, vacio: pd.DataFrame) -> pd.DataFrame:
    if not path.exists():
        return vacio.copy()
    try:
        df = pd.read_excel(path, sheet_name=hoja, dtype={"SKU": str})
        for c in vacio.columns:
            if c not in df.columns:
                df[c] = pd.NA
        return df[list(vacio.columns)].copy()
    except Exception as e:  # noqa: BLE001
        st.warning(f"No se pudo leer {path.name}: {e}. Se muestra tabla vacía.")
        return vacio.copy()


def _guardar_excel(df: pd.DataFrame, path: Path, hoja: str) -> None:
    DATA_DIR.mkdir(exist_ok=True)
    temp_path = path.with_name(f"{path.stem}.tmp_{os.getpid()}_{threading.get_ident()}.xlsx")
    try:
        with pd.ExcelWriter(temp_path, engine="openpyxl") as w:
            df.to_excel(w, sheet_name=hoja, index=False)
            ws = w.sheets[hoja]
            for cell in ws[1]:
                cell.fill = FILL_ENCABEZADO
                cell.font = FONT_ENCABEZADO
                cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.freeze_panes = "A2"
        # Reemplazo atómico para máxima seguridad ante cortes/fallos
        os.replace(temp_path, path)
    except Exception:
        if temp_path.exists():
            try:
                temp_path.unlink()
            except OSError:
                pass
        raise


@st.cache_data(ttl=60)
def cargar_inventario() -> pd.DataFrame:
    df = _leer_excel_seguro(INV_PATH, "Inventario", df_vacio_inv())
    if len(df):
        df["SKU"] = df["SKU"].astype(str).str.strip()
        # CORRECCIÓN 1: todo monetario y stock como enteros (sin .00)
        for c in ("Costo", "Precio_Sugerido", "Precio_Actual"):
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).round(0).astype(int)
        for c in ("Stock_Actual", "Stock_Minimo"):
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).round(0).astype(int)
        # CORRECCIÓN 4: recalcular SIEMPRE desde Costo
        df["Precio_Sugerido"] = (df["Costo"] * MARGEN).round(0).astype(int)
    return df


@st.cache_data(ttl=60)
def cargar_ventas() -> pd.DataFrame:
    df = _leer_excel_seguro(VENTAS_PATH, "Ventas", df_vacio_ventas())
    if len(df):
        df["Fecha"] = pd.to_datetime(df["Fecha"], errors="coerce")
        df["Cantidad"] = pd.to_numeric(df["Cantidad"], errors="coerce").fillna(0).round(0).astype(int)
        for c in ("Precio_Unitario", "Total"):
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).round(0).astype(int)
    return df


@st.cache_data(ttl=60)
def cargar_deudores() -> pd.DataFrame:
    df = _leer_excel_seguro(DEUD_PATH, "Deudores", df_vacio_deud())
    if len(df):
        for c in ("Monto_Total", "Monto_Pagado", "Saldo"):
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).round(0).astype(int)
    return df


def guardar_inventario(df: pd.DataFrame) -> None:
    df = df.copy()
    # CORRECCIÓN 4: la única fuente del sugerido es el Costo de Compra
    df["Costo"] = pd.to_numeric(df["Costo"], errors="coerce").fillna(0).round(0).astype(int)
    df["Precio_Actual"] = pd.to_numeric(df["Precio_Actual"], errors="coerce").fillna(0).round(0).astype(int)
    df["Stock_Actual"] = pd.to_numeric(df["Stock_Actual"], errors="coerce").fillna(0).round(0).astype(int)
    df["Stock_Minimo"] = pd.to_numeric(df["Stock_Minimo"], errors="coerce").fillna(0).round(0).astype(int)
    df["Precio_Sugerido"] = (df["Costo"] * MARGEN).round(0).astype(int)
    df["Fecha_Actualizacion"] = ahora_str()
    # Registrar categorías nuevas que aparezcan en el inventario
    for c in df["Categoria"].dropna().unique().tolist():
        guardar_nueva_categoria(str(c))
    _guardar_excel(df[COL_INV], INV_PATH, "Inventario")
    cargar_inventario.clear()


def guardar_ventas(df: pd.DataFrame) -> None:
    df = df.copy()
    df["Cantidad"] = pd.to_numeric(df["Cantidad"], errors="coerce").fillna(0).round(0).astype(int)
    for c in ("Precio_Unitario", "Total"):
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).round(0).astype(int)
    _guardar_excel(df[COL_VENTAS], VENTAS_PATH, "Ventas")
    cargar_ventas.clear()


def guardar_deudores(df: pd.DataFrame) -> None:
    df = df.copy()
    for c in ("Monto_Total", "Monto_Pagado", "Saldo"):
        df[c] = pd.to_numeric(df[c], errors="coerce").fillna(0).round(0).astype(int)
    _guardar_excel(df[COL_DEUD], DEUD_PATH, "Deudores")
    cargar_deudores.clear()


# ------------------------------------------------- Usuarios y autenticación
def hash_password(password: str) -> str:
    """Hash SHA-256 de la contraseña (no se guarda texto plano)."""
    return hashlib.sha256((password or "").encode("utf-8")).hexdigest()


def df_vacio_usu() -> pd.DataFrame:
    return pd.DataFrame(columns=COL_USU)


@st.cache_data(ttl=60)
def cargar_usuarios() -> pd.DataFrame:
    df = _leer_excel_seguro(USU_PATH, "Usuarios", df_vacio_usu())
    if len(df):
        df["Documento"] = df["Documento"].astype(str).str.strip()
        df["Activo"] = df["Activo"].astype(str).str.lower().isin(["true", "1", "sí", "si", "activo"])
    return df


def guardar_usuarios(df: pd.DataFrame) -> None:
    _guardar_excel(df[COL_USU], USU_PATH, "Usuarios")
    cargar_usuarios.clear()


def asegurar_superusuario() -> None:
    """Crea el SuperUsuario inicial si no existe (sin borrar los demás)."""
    df = cargar_usuarios()
    if len(df) and (df["Documento"].astype(str) == SUPER_DOC).any():
        return
    nuevo = pd.DataFrame([[SUPER_DOC, SUPER_NOMBRE, hash_password(SUPER_PASS),
                           "SuperUsuario", True, ahora_str()]], columns=COL_USU)
    guardar_usuarios(pd.concat([df, nuevo], ignore_index=True))


def rol_actual() -> str:
    auth = st.session_state.get("auth")
    return str(auth.get("Rol", "")) if auth else ""


def puede_editar_inventario() -> bool:
    """Solo SuperUsuario y Admin editan stock y datos críticos. Cajero: solo consulta."""
    return rol_actual() in ("SuperUsuario", "Admin")


def es_gestor_usuarios() -> bool:
    return rol_actual() in ("SuperUsuario", "Admin")


def super_activos(df: pd.DataFrame, excluir_doc: str = "") -> pd.DataFrame:
    """SuperUsuarios activos, opcionalmente excluyendo un documento.

    Robusto ante booleanos o textos en la columna Activo.
    """
    if not len(df):
        return df
    d = df[df["Documento"].astype(str) != str(excluir_doc)]
    activos = d["Activo"].astype(str).str.lower().isin(["true", "1", "sí", "si"])
    return d[(d["Rol"].astype(str) == "SuperUsuario") & activos]


def puede_crear_rol(rol_nuevo: str) -> bool:
    """Solo SuperUsuario puede crear/asignar el rol SuperUsuario."""
    if str(rol_nuevo) == "SuperUsuario" and rol_actual() != "SuperUsuario":
        return False
    return True


def cerrar_sesion() -> None:
    st.session_state.auth = None
    st.session_state.carrito = []
    st.session_state.items_deuda = []
    st.session_state.ultimo_ticket = None
    st.session_state.auto_print_ticket = False
    st.session_state.pagina = "📊 Panel Principal"
    st.rerun()


def estado_stock(stock: int, minimo: int) -> str:
    if minimo <= 0:
        return "✅ OK"
    if stock <= 0:
        return "🔴 Sin stock"
    if stock <= minimo:
        return "🔴 Crítico"
    if stock <= minimo * 1.5:
        return "🟡 Bajo"
    return "✅ OK"


def resaltar_stock(df: pd.DataFrame) -> pd.io.formats.style.Styler:
    """Colorea filas: rojo si Stock <= Mínimo, amarillo si <= Mínimo*1.5."""
    def color_fila(row):
        try:
            s, m = int(row["Stock_Actual"]), int(row["Stock_Minimo"])
        except (ValueError, TypeError, KeyError):
            return [""] * len(row)
        if s <= m:
            return ["background-color: #f8d7da; color: #58151c"] * len(row)  # rojo + texto oscuro legible
        if m > 0 and s <= m * 1.5:
            return ["background-color: #fff3cd; color: #5a4300"] * len(row)  # amarillo + texto oscuro
        return [""] * len(row)
    cols = [c for c in COL_INV if c in df.columns]
    return df[cols].style.apply(color_fila, axis=1).format(
        {c: lambda v: fmt_gs(v) if pd.notna(v) else "" for c in ("Costo", "Precio_Sugerido", "Precio_Actual")}
    )


def _colores_grafico() -> dict:
    """Paleta según el tema activo (claro/oscuro) para que los gráficos lean bien."""
    oscuro = False
    try:
        oscuro = st.get_option("theme.base") == "dark"
    except Exception:
        pass
    return {"texto": "#e8eaed" if oscuro else "#1f2937",
            "rejilla": "#3c4043" if oscuro else "#e5e7eb",
            "acento": "#4da3ff" if oscuro else "#1F4E79",
            "area": "#4da3ff" if oscuro else "#1F4E79"}


def _base_grafico(chart, titulo: str = ""):
    c = _colores_grafico()
    return (chart.properties(title=titulo)
            .configure(background="transparent")
            .configure_title(color=c["texto"], fontSize=15, anchor="start")
            .configure_axis(labelColor=c["texto"], titleColor=c["texto"],
                            gridColor=c["rejilla"], labelFontSize=11)
            .configure_legend(labelColor=c["texto"], titleColor=c["texto"])
            .configure_view(strokeWidth=0))


def graf_barras(datos: pd.DataFrame, etiqueta: str, valor: str, titulo: str = "",
                horizontal: bool = False, alto: int = 300):
    """Barras interactivas (hover con tooltip, resaltado) con formato de miles."""
    import altair as alt

    c = _colores_grafico()
    enc = ({"y": alt.Y(f"{etiqueta}:N", sort="-x", title=None),
            "x": alt.X(f"{valor}:Q", title="Total (Gs.)", axis=alt.Axis(format=","))}
           if horizontal else
           {"x": alt.X(f"{etiqueta}:N", sort="-y", title=None),
            "y": alt.Y(f"{valor}:Q", title="Total (Gs.)", axis=alt.Axis(format=","))})
    sel = alt.selection_point(fields=[etiqueta], on="click", clear="dblclick")
    chart = (alt.Chart(datos).mark_bar(cornerRadiusEnd=6, color=c["acento"])
             .encode(**enc, tooltip=[etiqueta, alt.Tooltip(valor, format=",", title="Gs.")],
                     opacity=alt.condition(sel, alt.value(1), alt.value(0.35)))
             .add_params(sel))
    return _base_grafico(chart.properties(height=alto), titulo)


def graf_dona(datos: pd.DataFrame, etiqueta: str, valor: str, titulo: str = "", alto: int = 300):
    """Dona interactiva con tooltip y resaltado al pasar el cursor."""
    import altair as alt

    c = _colores_grafico()
    sel = alt.selection_point(fields=[etiqueta], on="mouseover", clear="mouseout")
    chart = (alt.Chart(datos).mark_arc(innerRadius=70, outerRadius=120, cornerRadius=5)
             .encode(theta=alt.Theta(f"{valor}:Q"),
                     color=alt.Color(f"{etiqueta}:N",
                                     scale=alt.Scale(scheme="blues" if c["texto"] == "#1f2937" else "lighttealblue"),
                                     legend=alt.Legend(title=None)),
                     tooltip=[etiqueta, alt.Tooltip(valor, format=",", title="Gs.")],
                     opacity=alt.condition(sel, alt.value(1), alt.value(0.45)))
             .add_params(sel))
    return _base_grafico(chart.properties(height=alto), titulo)


def graf_serie(datos: pd.DataFrame, x: str, valor: str, titulo: str = "", alto: int = 320):
    """Serie temporal con área, línea y puntos interactivos."""
    import altair as alt

    c = _colores_grafico()
    base = alt.Chart(datos).encode(
        x=alt.X(f"{x}:T", title=None),
        y=alt.Y(f"{valor}:Q", title="Total (Gs.)", axis=alt.Axis(format=",")),
        tooltip=[alt.Tooltip(f"{x}:T", title="Fecha"),
                 alt.Tooltip(valor, format=",", title="Gs.")])
    area = base.mark_area(opacity=0.25, color=c["area"])
    linea = base.mark_line(color=c["acento"], strokeWidth=2.5)
    puntos = base.mark_circle(size=70, color=c["acento"]).encode(
        opacity=alt.condition(alt.selection_point(on="mouseover"), alt.value(1), alt.value(0.85)))
    chart = (area + linea + puntos).interactive()
    return _base_grafico(chart.properties(height=alto), titulo)


def df_a_excel_bytes(hojas: dict[str, pd.DataFrame], con_totales: bool = True) -> bytes:
    """Exporta planillas .xlsx con formato profesional.

    - Cabeceras con fondo azul y texto blanco en negrita.
    - Ancho de columnas autoajustado.
    - Números sin decimales; moneda en formato Gs. (Excel: "Gs." #,##0).
    - Fila de TOTALES al final (suma de cantidades y montos) + filtros y panel congelado.
    """
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for nombre, df in hojas.items():
            hoja = (nombre[:31] or "Hoja")
            datos = df.copy()
            if con_totales and len(datos):
                datos = _agregar_fila_totales(datos)
            datos.to_excel(w, sheet_name=hoja, index=False)
            _formato_profesional(w.sheets[hoja], datos)
    buf.seek(0)
    return buf.getvalue()


COLS_MONEDA = {"Costo", "Precio_Sugerido", "Precio_Actual", "Precio_Unitario",
               "Total", "Subtotal", "Monto_Total", "Monto_Pagado", "Saldo", "Ganancia"}
COLS_CANTIDAD = {"Cantidad", "Stock_Actual", "Stock_Minimo", "Unidades"}
FMT_MONEDA_XL = '"Gs." #,##0'
FMT_CANT_XL = '#,##0'
FILL_ENCABEZADO = PatternFill("solid", fgColor="1F4E79")
FONT_ENCABEZADO = Font(bold=True, color="FFFFFF", size=11)
FILL_TOTAL = PatternFill("solid", fgColor="DDEBF7")
FONT_TOTAL = Font(bold=True, size=11)
BORDE_FINO = Border(left=Side(style="thin", color="BFBFBF"),
                    right=Side(style="thin", color="BFBFBF"),
                    top=Side(style="thin", color="BFBFBF"),
                    bottom=Side(style="thin", color="BFBFBF"))


def _agregar_fila_totales(df: pd.DataFrame) -> pd.DataFrame:
    """Añade fila 'TOTAL (N filas)' con la suma de columnas de cantidad/moneda."""
    total: dict[str, object] = {}
    for c in df.columns:
        if c in COLS_MONEDA or c in COLS_CANTIDAD:
            total[c] = int(pd.to_numeric(df[c], errors="coerce").fillna(0).sum())
        else:
            total[c] = ""
    total[df.columns[0]] = f"TOTAL ({len(df)} filas)"
    return pd.concat([df, pd.DataFrame([total])], ignore_index=True)


def _formato_profesional(ws, df: pd.DataFrame) -> None:
    ncols, nrows = df.shape[0] + 1, df.shape[1]
    # Cabecera
    for col_idx in range(1, nrows + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = FILL_ENCABEZADO
        cell.font = FONT_ENCABEZADO
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDE_FINO
    ws.row_dimensions[1].height = 24
    # Fila de totales
    if len(df) and str(ws.cell(row=ncols, column=1).value or "").startswith("TOTAL"):
        for col_idx in range(1, nrows + 1):
            cell = ws.cell(row=ncols, column=col_idx)
            cell.fill = FILL_TOTAL
            cell.font = FONT_TOTAL
            cell.border = BORDE_FINO
        ws.row_dimensions[ncols].height = 20
    # Formatos numéricos + bordes del cuerpo
    for col_idx, col in enumerate(df.columns, start=1):
        for row_idx in range(2, ncols + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = BORDE_FINO
            if col in COLS_MONEDA:
                cell.number_format = FMT_MONEDA_XL
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col in COLS_CANTIDAD:
                cell.number_format = FMT_CANT_XL
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif row_idx > 1:
                cell.alignment = Alignment(vertical="center", wrap_text=True)
    # Ancho autoajustado (seguro ante nulos pd.NA/NaN: se usa pd.isna)
    for col_idx, col in enumerate(df.columns, start=1):
        letra = get_column_letter(col_idx)
        max_len = len(str(col))
        for val in list(df[col].head(300)) + [ws.cell(row=ncols, column=col_idx).value if len(df) else ""]:
            if val is None:
                continue
            try:
                if bool(pd.isna(val)):
                    continue
            except (TypeError, ValueError):
                pass
            if isinstance(val, str) and val == "":
                continue
            if col in COLS_MONEDA:
                try:
                    largo = len(f"Gs. {int(val):,}".replace(",", ".")) + 2
                except (TypeError, ValueError):
                    largo = len(str(val))
            else:
                largo = len(str(val))
            max_len = max(max_len, min(largo, 60))
        ws.column_dimensions[letra].width = min(max(max_len + 2, 13), 52)
    ws.freeze_panes = "A2"
    if ncols > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(nrows)}{ncols - (1 if len(df) and str(ws.cell(row=ncols, column=1).value or '').startswith('TOTAL') else 0)}"


def ticket_a_html_doc(t: dict, auto_print: bool = False) -> str:
    """Genera el documento HTML del ticket de venta (optimizado para impresora de 80mm).

    Se imprime con window.print() dentro de su iframe: solo sale el ticket,
    y el botón de impresión se oculta automáticamente en papel (@media print).
    """
    filas = "\n".join(
        "<tr><td>" + html.escape(str(it["Nombre"])) + "</td>"
        + f"<td class='c'>{int(it['Cantidad'])}</td>"
        + f"<td class='r'>{fmt_gs(it['Precio'])}</td>"
        + f"<td class='r'>{fmt_gs(int(it['Cantidad']) * a_entero(it['Precio']))}</td></tr>"
        for it in t["items"])
    auto = "<script>window.onload=function(){window.print();};</script>" if auto_print else ""
    notas = f"<p class='nota'>Nota: {html.escape(str(t.get('notas') or ''))}</p>" if t.get("notas") else ""
    vend = (f"<p>Atendido por: <b>{html.escape(str(t.get('vendedor') or ''))}</b></p>"
            if t.get("vendedor") else "")
    return f"""<html><head><meta charset="utf-8"><style>
      body{{font-family:'Courier New',Courier,monospace;margin:0;padding:8px;color:#000;}}
      .ticket{{max-width:320px;margin:0 auto;}}
      h2,p{{text-align:center;margin:4px 0;font-size:14px;}}
      h2{{font-size:18px;}}
      table{{width:100%;border-collapse:collapse;font-size:13px;}}
      th{{border-top:1px dashed #000;border-bottom:1px dashed #000;text-align:left;padding:4px 2px;}}
      td{{padding:3px 2px;border-bottom:1px dotted #bbb;vertical-align:top;}}
      .c{{text-align:center;}}.r{{text-align:right;white-space:nowrap;}}
      .total{{font-size:18px;font-weight:bold;text-align:right;border-top:2px solid #000;margin-top:6px;padding-top:6px;}}
      .nota{{font-size:12px;}}
      .btn{{display:block;width:100%;margin:10px 0;padding:10px;font-size:15px;font-weight:bold;
            background:#1F4E79;color:#fff;border:none;border-radius:6px;cursor:pointer;}}
      @media print{{.btn{{display:none;}}body{{padding:0;}}}}
      </style></head><body><div class="ticket">
      <h2>📚 Librería MM</h2>
      <p>Ticket <b>{html.escape(str(t['id']))}</b><br>{html.escape(str(t['fecha']))} · {html.escape(str(t['caja']))}</p>
      <p>Método de pago: <b>{html.escape(str(t['metodo']))}</b></p>
      {vend}
      <table><tr><th>Producto</th><th>Cant.</th><th>P.Unit</th><th>Subt.</th></tr>{filas}</table>
      <div class="total">TOTAL: {fmt_gs(t['total'])}</div>
      {notas}
      <p>¡Gracias por su compra!</p>
      <button class="btn" onclick="window.print()">🖨️ Imprimir ticket</button>
      </div>{auto}</body></html>"""


# ------------------------------------------- Importación flexible (Excel viejo)
SINONIMOS = {
    "sku": "SKU", "codigo": "SKU", "código": "SKU", "cod": "SKU",
    "codigo_barras": "SKU", "cod_barras": "SKU", "codigobarra": "SKU", "id": "SKU",
    "nombre": "Nombre", "producto": "Nombre", "descripcion": "Nombre",
    "descripción": "Nombre", "detalle": "Nombre", "articulo": "Nombre", "artículo": "Nombre",
    "categoria": "Categoria", "categoría": "Categoria", "rubro": "Categoria", "tipo": "Categoria",
    "costo": "Costo", "precio_compra": "Costo", "precio_costo": "Costo", "compra": "Costo",
    "costo_compra": "Costo", "costo_adquisicion": "Costo",
    "precio_sugerido": "Precio_Sugerido", "sugerido": "Precio_Sugerido",
    "precio": "Precio_Actual", "precio_venta": "Precio_Actual", "precio_actual": "Precio_Actual",
    "pvp": "Precio_Actual", "venta": "Precio_Actual",
    "stock": "Stock_Actual", "stock_actual": "Stock_Actual", "cantidad": "Stock_Actual",
    "existencia": "Stock_Actual", "disponible": "Stock_Actual",
    "minimo": "Stock_Minimo", "mínimo": "Stock_Minimo", "stock_minimo": "Stock_Minimo",
    "min": "Stock_Minimo", "alerta": "Stock_Minimo", "critico": "Stock_Minimo",
    "fecha": "Fecha", "dia": "Fecha", "día": "Fecha",
    "metodo_pago": "Metodo_Pago", "pago": "Metodo_Pago", "medio_pago": "Metodo_Pago",
    "forma_pago": "Metodo_Pago", "metodo": "Metodo_Pago",
    "cliente": "Cliente", "nombre_cliente": "Cliente", "deudor": "Cliente",
    "monto": "Monto_Total", "total": "Total", "deuda": "Monto_Total",
    "importe": "Monto_Total", "debe": "Saldo", "saldo": "Saldo",
}

def _norm(col: str) -> str:
    c = str(col).strip().lower()
    c = re.sub(r"[^a-z0-9áéíóúñ]+", "_", c).strip("_")
    c = (c.replace("á", "a").replace("é", "e").replace("í", "i")
           .replace("ó", "o").replace("ú", "u"))
    return SINONIMOS.get(c, str(col).strip())


def normalizar_df_inventario(crudo: pd.DataFrame) -> pd.DataFrame:
    df = crudo.rename(columns={c: _norm(c) for c in crudo.columns}).copy()
    for c in COL_INV:
        if c not in df.columns:
            df[c] = pd.NA
    df["SKU"] = df["SKU"].astype(str).str.strip()
    df = df[df["SKU"].notna() & (df["SKU"] != "") & (df["SKU"].str.lower() != "nan")]
    df["Nombre"] = df["Nombre"].fillna("Sin descripción").astype(str)
    df["Categoria"] = df["Categoria"].fillna("Otros").astype(str)
    df["Costo"] = pd.to_numeric(df["Costo"], errors="coerce").fillna(0).round(0).astype(int)
    df["Precio_Actual"] = pd.to_numeric(df["Precio_Actual"], errors="coerce")
    # CORRECCIÓN 4: sugerido SIEMPRE desde Costo de Compra
    df["Precio_Sugerido"] = (df["Costo"] * MARGEN).round(0).astype(int)
    df["Precio_Actual"] = df["Precio_Actual"].fillna(df["Precio_Sugerido"]).round(0).astype(int)
    df["Stock_Actual"] = pd.to_numeric(df["Stock_Actual"], errors="coerce").fillna(0).round(0).astype(int)
    df["Stock_Minimo"] = pd.to_numeric(df["Stock_Minimo"], errors="coerce").fillna(0).round(0).astype(int)
    df["Fecha_Actualizacion"] = ahora_str()
    return df[COL_INV].drop_duplicates(subset="SKU", keep="last").reset_index(drop=True)


def normalizar_df_ventas(crudo: pd.DataFrame, inv: pd.DataFrame) -> pd.DataFrame:
    df = crudo.rename(columns={c: _norm(c) for c in crudo.columns}).copy()
    if "Nombre" in df.columns and "Producto" not in df.columns:
        df = df.rename(columns={"Nombre": "Producto"})
    if "Precio_Actual" in df.columns and "Precio_Unitario" not in df.columns:
        df = df.rename(columns={"Precio_Actual": "Precio_Unitario"})
    precios = {}
    if len(inv):
        precios = dict(zip(inv["SKU"].astype(str), zip(inv["Nombre"], inv["Categoria"], inv["Precio_Actual"])))
    regs = []
    for i, r in df.iterrows():
        sku = str(r.get("SKU", "") or "").strip()
        prod = str(r.get("Producto", "") or "").strip() or "Ítem importado"
        cat, pu = "Otros", 0
        if sku in precios:
            prod, cat, pu = precios[sku][0], precios[sku][1], a_entero(precios[sku][2])
        elif prod != "Ítem importado" and len(inv):
            match = inv[inv["Nombre"].str.lower() == prod.lower()]
            if len(match):
                m = match.iloc[0]
                sku, cat, pu = str(m["SKU"]), str(m["Categoria"]), a_entero(m["Precio_Actual"])
        cant = a_entero(pd.to_numeric(pd.Series([r.get("Cantidad", 1)]), errors="coerce").fillna(1).iloc[0])
        if cant <= 0:
            cant = 1
        pu = a_entero(pd.to_numeric(pd.Series([r.get("Precio_Unitario", pu)]), errors="coerce").fillna(pu).iloc[0])
        fecha = pd.to_datetime(r.get("Fecha", datetime.now()), errors="coerce")
        if pd.isna(fecha):
            fecha = datetime.now()
        total = int(cant * pu)
        regs.append([f"V-IMP-{fecha.strftime('%Y%m%d')}-{i+1:04d}", fecha,
                     sku or "IMPORTADO", prod, str(r.get("Categoria", cat)),
                     int(cant), int(pu), int(total),
                     str(r.get("Metodo_Pago", "Efectivo")), str(r.get("Caja", "Caja 1")),
                     str(r.get("Notas", "Importado desde Excel")), "", ""])
    out = pd.DataFrame(regs, columns=COL_VENTAS)
    out["Fecha"] = pd.to_datetime(out["Fecha"])
    return out


def normalizar_df_deudores(crudo: pd.DataFrame) -> pd.DataFrame:
    df = crudo.rename(columns={c: _norm(c) for c in crudo.columns}).copy()
    for c in COL_DEUD:
        if c not in df.columns:
            df[c] = pd.NA
    regs = []
    for i, r in df.iterrows():
        tot = a_entero(pd.to_numeric(pd.Series([r.get("Monto_Total", r.get("Total", 0))]), errors="coerce").fillna(0).iloc[0])
        pag = a_entero(pd.to_numeric(pd.Series([r.get("Monto_Pagado", 0)]), errors="coerce").fillna(0).iloc[0])
        regs.append([str(r.get("ID", f"D-IMP-{i+1:04d}")), str(r.get("Cliente", "Sin nombre")),
                     str(r.get("Telefono", "")), str(r.get("Concepto", "Deuda importada")),
                     int(tot), int(pag), int(tot - pag),
                     str(r.get("Fecha_Prestamo", date.today())),
                     str(r.get("Fecha_Vencimiento", date.today())),
                     str(r.get("Estado", "Pendiente" if tot - pag > 0 else "Pagado"))])
    return pd.DataFrame(regs, columns=COL_DEUD)


# ------------------------------------------------------------------ Estado
if "carrito" not in st.session_state:
    st.session_state.carrito = []  # [{SKU, Nombre, Cantidad, Precio}]
if "items_deuda" not in st.session_state:
    st.session_state.items_deuda = []  # [{Producto, Cantidad, Precio, Subtotal}]
if "ultimo_ticket" not in st.session_state:
    st.session_state.ultimo_ticket = None  # dict del ticket recién cobrado
if "auto_print_ticket" not in st.session_state:
    st.session_state.auto_print_ticket = False  # dispara window.print() una vez
if "auth" not in st.session_state:
    st.session_state.auth = None  # {"Documento","Nombre","Rol"} del usuario logueado
if "pagina" not in st.session_state:
    st.session_state.pagina = "📊 Panel Principal"
if st.session_state.pagina == "📊 Dashboard":  # migración desde el nombre anterior
    st.session_state.pagina = "📊 Panel Principal"

# Sincronización inicial con PRIORIDAD A LA NUBE: al abrir, la interfaz refleja
# los datos oficiales de Drive (lo local obsoleto se respalda, jamás sube).
# Se ejecuta COMO MÁXIMO 1 vez cada 10 min (típicamente 1 por sesión): hacerla
# en cada rerun colgaba cada clic ("Agregar al ticket", cambio de menú) con
# llamadas de red a Drive. El resto de la navegación trabaja 100% en local.
_sync_ahora = time.time()
if _sync_ahora - float(st.session_state.get("_sync_inicio_ts", 0)) > 600:
    drive_sync.sincronizar_inicio()
    st.session_state["_sync_inicio_ts"] = _sync_ahora
    st.cache_data.clear()  # lo recién descargado debe verse de inmediato
inv = cargar_inventario()
ventas = cargar_ventas()
deud = cargar_deudores()
usuarios = cargar_usuarios()
asegurar_superusuario()

# ------------------------------------------------------- Puerta de seguridad
if not st.session_state.auth:
    c_izq, c_login, c_der = st.columns([1, 2, 1])
    with c_login:
        st.title("📚 Librería MM")
        st.subheader("🔐 Iniciar sesión")
        st.caption("Ingresa tu documento y contraseña para continuar.")
        with st.form("login", clear_on_submit=False):
            doc = st.text_input("Número de Documento (Cédula)")
            pwd = st.text_input("Contraseña", type="password")
            if st.form_submit_button("➡️ Ingresar", type="primary", use_container_width=True):
                dfu = cargar_usuarios()
                ok = False
                if len(dfu):
                    m = dfu[dfu["Documento"].astype(str).str.strip() == doc.strip()]
                    if len(m):
                        u = m.iloc[0]
                        if bool(u["Activo"]) and str(u["Password_Hash"]) == hash_password(pwd):
                            st.session_state.auth = {
                                "Documento": str(u["Documento"]),
                                "Nombre": str(u["Nombre"]), "Rol": str(u["Rol"])}
                            ok = True
                if ok:
                    st.rerun()
                else:
                    st.error("Documento, contraseña o usuario inactivo incorrectos.")
    st.stop()

# Carga de inventario base (necesario para badge de alertas en sidebar y categorías)
# (ventas/deud/usuarios ya se cargaron arriba, antes del login).
CATEGORIAS = obtener_categorias(inv)

# ------------------------------------------------------------------ Sidebar
MENU_COMPLETO = ["📊 Panel Principal", "📦 Inventario", "🧾 Ventas y Caja",
                 "🤝 Deudores", "🔄 Importar / Exportar", "📈 Reportes", "👥 Usuarios"]
MENU_CAJERO = ["📊 Panel Principal", "📦 Inventario", "🧾 Ventas y Caja"]
MENU = MENU_CAJERO if rol_actual() == "Cajero" else MENU_COMPLETO
if st.session_state.pagina not in MENU:
    st.session_state.pagina = MENU[0]
with st.sidebar:
    st.title("📚 Librería MM")
    st.caption("Inventario · Caja · Deudores · Gs.")
    st.caption(f"Almacenamiento: **{drive_sync.estado_almacenamiento()}**")
    auth = st.session_state.get("auth") or {}
    if auth:
        st.info(f"👤 **{auth.get('Nombre', '')}**\n\n🪪 {auth.get('Documento', '')} · 🎭 {auth.get('Rol', '')}")
    if st.button("🚪 Cerrar Sesión", use_container_width=True):
        cerrar_sesion()
    if st.button("🔄 Sincronizar", use_container_width=True,
                 help="Descarga la versión oficial de Drive (prioridad a la nube); lo local se respalda en data/_backup_sync/ antes de sobrescribir"):
        with st.spinner("Sincronizando con Drive..."):
            ok_sync, det_sync = drive_sync.sincronizacion_manual()
        st.session_state["_sync_msg"] = (ok_sync, det_sync)
        st.cache_data.clear()
        st.rerun()
    st.divider()
    st.session_state.pagina = st.radio("Menú", MENU, index=MENU.index(st.session_state.pagina))
    st.divider()
    n_crit = int(((inv["Stock_Actual"] <= inv["Stock_Minimo"]) & (len(inv) > 0)).sum()) if len(inv) else 0
    st.metric("⚠️ Productos en alerta", n_crit)
    if not INV_PATH.exists():
        st.warning("Aún no hay base local. Ejecuta `python init_db.py` o carga productos en Inventario.")

pg = st.session_state.pagina

_sync_msg = st.session_state.pop("_sync_msg", None)
if _sync_msg:
    _ok_m, _det_m = _sync_msg
    if _ok_m:
        st.success(f"✅ {_det_m}")
    else:
        st.error(f"❌ {_det_m}")

# --------------------------------------------------------------- PANEL PRINCIPAL
if pg == "📊 Panel Principal":
    st.header("📊 Panel general")
    hoy = date.today()
    v = ventas.copy()
    if len(v):
        v["Fecha"] = pd.to_datetime(v["Fecha"], errors="coerce")
        v = v.dropna(subset=["Fecha"])
    v_hoy = v[v["Fecha"].dt.date == hoy] if len(v) else v
    v_mes = v[(v["Fecha"].dt.year == hoy.year) & (v["Fecha"].dt.month == hoy.month)] if len(v) else v

    c1, c2, c3, c4 = st.columns(4)
    c1.metric("💰 Ventas hoy", fmt_gs(v_hoy["Total"].sum()) if len(v_hoy) else fmt_gs(0), f"{len(v_hoy)} ítems")
    c2.metric("🗓️ Ventas del mes", fmt_gs(v_mes["Total"].sum()) if len(v_mes) else fmt_gs(0))
    c3.metric("📦 Productos", f"{len(inv):,}".replace(",", "."))
    if rol_actual() == "Cajero":
        # El cajero no ve costos: se muestra conteo de alertas en su lugar
        n_bajo = int(((inv["Stock_Actual"] <= inv["Stock_Minimo"]) & (len(inv) > 0)).sum()) if len(inv) else 0
        c4.metric("⚠️ En alerta", f"{n_bajo:,}".replace(",", "."))
    else:
        valor_costo = int((inv["Costo"] * inv["Stock_Actual"]).sum()) if len(inv) else 0
        c4.metric("🏷️ Valor inventario (costo)", fmt_gs(valor_costo))

    col_a, col_b = st.columns([2, 1])
    with col_a:
        st.subheader("Ventas por día (últimos 14 días)")
        if len(v):
            ult = v[v["Fecha"] >= pd.Timestamp(hoy) - pd.Timedelta(days=13)]
            if len(ult):
                diario = ult.groupby(ult["Fecha"].dt.date)["Total"].sum().reset_index()
                diario.columns = ["Fecha", "Total"]
                diario["Fecha"] = pd.to_datetime(diario["Fecha"])
                st.altair_chart(graf_serie(diario, "Fecha", "Total"), use_container_width=True)
            else:
                st.info("Sin ventas en los últimos 14 días.")
        else:
            st.info("Todavía no hay ventas registradas. Ve a 🧾 Ventas y Caja.")
    with col_b:
        st.subheader("Ventas por método de pago")
        if len(v):
            met = v.groupby("Metodo_Pago", as_index=False)["Total"].sum()
            st.altair_chart(graf_dona(met, "Metodo_Pago", "Total"), use_container_width=True)
        st.subheader("Top categorías")
        if len(v):
            cat = (v.groupby("Categoria", as_index=False)["Total"].sum()
                   .sort_values("Total", ascending=False).head(5))
            st.altair_chart(graf_barras(cat, "Categoria", "Total", horizontal=True, alto=220),
                            use_container_width=True)

    st.subheader("🚨 Alertas de stock (Stock ≤ Mínimo)")
    if len(inv):
        alertas = inv[inv["Stock_Actual"] <= inv["Stock_Minimo"]].copy()
        if len(alertas):
            alertas["Estado"] = alertas.apply(lambda r: estado_stock(r["Stock_Actual"], r["Stock_Minimo"]), axis=1)
            st.dataframe(resaltar_stock(alertas), use_container_width=True)
            da1, da2 = st.columns(2)
            da1.download_button("⬇️ Descargar alertas (.xlsx)",
                                df_a_excel_bytes({"Alertas": alertas[COL_INV]}),
                                file_name=f"alertas-MM-{date.today()}.xlsx",
                                key="dl_alert_xlsx")
            da2.download_button("⬇️ Descargar alertas (.csv con ;)",
                                csv_puntoycoma_bytes(alertas[COL_INV]),
                                file_name=f"alertas-MM-{date.today()}.csv",
                                mime="text/csv",
                                key="dl_alert_csv")
        else:
            st.success("✅ Sin alertas: todo el stock está por encima del mínimo.")
    else:
        st.info("Carga productos para activar las alertas.")

# -------------------------------------------------------------- INVENTARIO
elif pg == "📦 Inventario":
    st.header("📦 Inventario (Gs.)")
    if not puede_editar_inventario():
        # Rol Cajero: solo consulta de stock y precios de venta (sin costos ni edición)
        st.info("🔒 Tu rol (**Cajero**) es de solo consulta: puedes ver stock y precios de venta, "
                "pero no modificar existencias, costos ni categorías.")
        rc1, rc2 = st.columns([2, 1])
        busq = rc1.text_input("🔎 Buscar (SKU o nombre)", "", key="inv_ro_busq")
        fcat = rc2.selectbox("Categoría", ["Todas"] + CATEGORIAS, key="inv_ro_cat")
        vista = inv.copy()
        if busq:
            vista = vista[vista["SKU"].str.contains(busq, case=False, na=False)
                          | vista["Nombre"].str.contains(busq, case=False, na=False)]
        if fcat != "Todas":
            vista = vista[vista["Categoria"] == fcat]
        if len(vista):
            vista["Estado"] = vista.apply(lambda r: estado_stock(r["Stock_Actual"], r["Stock_Minimo"]), axis=1)
            ver = vista[["SKU", "Nombre", "Categoria", "Precio_Actual", "Stock_Actual", "Estado"]].copy()
            ver["Precio_Actual"] = ver["Precio_Actual"].map(fmt_gs)
            st.dataframe(ver, use_container_width=True)
        else:
            st.info("Sin productos para mostrar.")
        st.stop()
    t1, t2, t3, t4 = st.tabs(["📋 Catálogo (editar)", "➕ Agregar producto", "✏️ Editar producto", "🏷️ Categorías"])

    with t4:
        # CORRECCIÓN 2: gestión de categorías como elemento interactivo real
        st.subheader("🏷️ Gestionar categorías")
        st.caption("Añade categorías nuevas aquí. Se guardan en `data/categorias.txt` y aparecen "
                   "en el selector de productos, filtros y reportes.")
        cols = st.columns([3, 1])
        nueva = cols[0].text_input("Nueva categoría", placeholder="Ej. Textos escolares",
                                   key="nueva_cat_global")
        if cols[1].button("➕ Añadir", type="primary", use_container_width=True):
            if guardar_nueva_categoria(nueva):
                st.success(f"✅ Categoría «{nueva.strip()}» añadida.")
                cargar_inventario.clear()
                st.rerun()
            else:
                st.warning("Escribe un nombre válido y no repetido.")
        st.write("**Categorías actuales:**")
        st.write(" · ".join(f"`{c}`" for c in CATEGORIAS) if CATEGORIAS else "Sin categorías.")
        st.divider()
        st.subheader("🗑️ Borrar categoría (borrado seguro)")
        st.caption("Los productos que usaban la categoría borrada se reasignan "
                   "automáticamente a «Otros». «Otros» no se puede borrar.")
        borrables = [c for c in CATEGORIAS if c.lower() != CATEGORIA_DEFECTO.lower()]
        if not borrables:
            st.info("No hay categorías borrables.")
        else:
            cat_del = st.selectbox("Categoría a borrar", borrables, key="cat_del_sel")
            n_af = int((inv["Categoria"].astype(str).str.strip().str.lower() == cat_del.lower()).sum()) \
                if len(inv) else 0
            st.warning(f"«{cat_del}» está en {n_af} producto(s). Al borrarla pasarán a «Otros».")
            conf = st.checkbox(f"Confirmo borrar «{cat_del}»", key="cat_del_conf")
            if st.button("🗑️ Borrar categoría", key="cat_del_go"):
                if not conf:
                    st.warning("Marca la confirmación primero.")
                else:
                    ok, n = eliminar_categoria(cat_del)
                    if ok:
                        st.success(f"✅ Categoría «{cat_del}» borrada. {n} producto(s) reasignados a «Otros».")
                        st.rerun()
                    else:
                        st.error("No se pudo borrar («Otros» está protegida).")

    with t2:
        # CORRECCIÓN 2: selector funcional + input dinámico que guarda la categoría
        st.subheader("Nuevo producto")
        CATEGORIAS = obtener_categorias(inv)
        with st.form("nuevo_prod", clear_on_submit=True):
            c1, c2 = st.columns(2)
            sku = c1.text_input("SKU / Código de barras *")
            nombre = c2.text_input("Nombre del producto *")
            st.markdown("**Categoría ***")
            cc1, cc2 = st.columns([1, 1])
            cat_sel = cc1.selectbox("Elige de la lista", CATEGORIAS, key="cat_sel")
            cat_txt = cc2.text_input("…o escribe una nueva", placeholder="Ej. Librería fina",
                                     key="cat_txt")
            st.caption("Si escribes una categoría nueva, se añade automáticamente a la lista y queda guardada.")
            c4, c5, c6, c7 = st.columns(4)
            costo = c4.number_input("Costo de Compra (Gs.) *", min_value=0, step=500, format="%d")
            precio = c5.number_input("Precio Actual de venta (Gs.)", min_value=0, step=500, format="%d")
            stock = c6.number_input("Stock Actual", min_value=0, step=1, format="%d")
            minimo = c7.number_input("Stock Mínimo (alerta)", min_value=0, step=1, format="%d")
            # CORRECCIÓN 4: vista previa calculada desde Costo
            st.info(f"Precio Sugerido automático = Costo × 1.5 → **{fmt_gs(recalcular_sugerido(costo))}** "
                    f"(calculado desde el Costo de Compra, no desde la venta)")
            if st.form_submit_button("💾 Guardar producto", type="primary"):
                categoria = cat_txt.strip() or cat_sel
                if not sku.strip() or not nombre.strip():
                    st.error("SKU y Nombre son obligatorios.")
                elif sku.strip() in set(inv["SKU"].astype(str)):
                    st.error(f"El SKU `{sku}` ya existe. Usa otro o edítalo en el catálogo.")
                elif a_entero(costo) <= 0:
                    st.error("El Costo de Compra debe ser mayor a 0 (es la base del sugerido).")
                else:
                    if a_entero(precio) == 0:
                        precio = recalcular_sugerido(costo)
                    guardar_nueva_categoria(categoria)
                    nuevo = pd.DataFrame([[sku.strip(), nombre.strip(), categoria, a_entero(costo),
                                           recalcular_sugerido(costo), a_entero(precio),
                                           int(stock), int(minimo), ahora_str()]], columns=COL_INV)
                    guardar_inventario(pd.concat([inv, nuevo], ignore_index=True))
                    st.success(f"✅ `{nombre}` guardado. Sugerido: {fmt_gs(recalcular_sugerido(costo))}")
                    st.rerun()

    with t3:
        st.subheader("✏️ Editar producto ya creado")
        if not len(inv):
            st.info("No hay productos para editar. Agrega uno en ➕ Agregar producto.")
        else:
            CATEGORIAS = obtener_categorias(inv)
            sel_sku = st.selectbox(
                "Elige el producto",
                [f"{r.SKU} — {r.Nombre}" for r in inv.itertuples()],
                key="edit_prod_sel")
            sku_e = sel_sku.split(" — ")[0]
            fila_e = inv[inv["SKU"] == sku_e].iloc[0]
            st.caption(f"Sugerido actual (Costo × 1.5): **{fmt_gs(fila_e['Precio_Sugerido'])}** · "
                       f"Última actualización: {fila_e['Fecha_Actualizacion']}")
            ec1, ec2 = st.columns(2)
            e_nom = ec1.text_input("Nombre", value=str(fila_e["Nombre"]), key=f"ep_nom_{sku_e}")
            e_cat_lista = ec2.selectbox("Categoría (lista)", CATEGORIAS,
                                        index=CATEGORIAS.index(str(fila_e["Categoria"]))
                                        if str(fila_e["Categoria"]) in CATEGORIAS else 0,
                                        key=f"ep_cat_{sku_e}")
            e_cat_txt = ec2.text_input("…o escribe una nueva categoría", key=f"ep_catnew_{sku_e}")
            e1, e2, e3, e4 = st.columns(4)
            e_costo = e1.number_input("Costo de Compra (Gs.)", min_value=0, step=500,
                                      value=a_entero(fila_e["Costo"]), format="%d", key=f"ep_cos_{sku_e}")
            e_precio = e2.number_input("Precio Actual de venta (Gs.)", min_value=0, step=500,
                                       value=a_entero(fila_e["Precio_Actual"]), format="%d", key=f"ep_pre_{sku_e}")
            e_stock = e3.number_input("Stock Actual", min_value=0, step=1,
                                      value=int(fila_e["Stock_Actual"]), format="%d", key=f"ep_stk_{sku_e}")
            e_min = e4.number_input("Stock Mínimo", min_value=0, step=1,
                                    value=int(fila_e["Stock_Minimo"]), format="%d", key=f"ep_min_{sku_e}")
            st.info(f"Nuevo Sugerido = Costo × 1.5 → **{fmt_gs(recalcular_sugerido(e_costo))}**")
            if st.button("💾 Guardar edición", type="primary", key=f"ep_save_{sku_e}"):
                if not e_nom.strip():
                    st.error("El Nombre no puede quedar vacío.")
                elif a_entero(e_costo) <= 0:
                    st.error("El Costo de Compra debe ser mayor a 0.")
                else:
                    categoria = e_cat_txt.strip() or e_cat_lista
                    guardar_nueva_categoria(categoria)
                    base = inv.set_index("SKU")
                    base.loc[sku_e, ["Nombre", "Categoria", "Costo", "Precio_Actual",
                                     "Stock_Actual", "Stock_Minimo"]] = [
                        e_nom.strip(), categoria, a_entero(e_costo), a_entero(e_precio),
                        int(e_stock), int(e_min)]
                    guardar_inventario(base.reset_index())
                    st.success(f"✅ `{sku_e}` actualizado. Nuevo sugerido: {fmt_gs(recalcular_sugerido(e_costo))}")
                    st.rerun()

    with t1:
        CATEGORIAS = obtener_categorias(inv)
        c1, c2, c3 = st.columns([2, 1, 1])
        busq = c1.text_input("🔎 Buscar (SKU o nombre)", "")
        fcat = c2.selectbox("Categoría", ["Todas"] + CATEGORIAS)
        solo_alertas = c3.checkbox("Solo alertas ⚠️")
        vista = inv.copy()
        if busq:
            vista = vista[vista["SKU"].str.contains(busq, case=False, na=False)
                          | vista["Nombre"].str.contains(busq, case=False, na=False)]
        if fcat != "Todas":
            vista = vista[vista["Categoria"] == fcat]
        if solo_alertas:
            vista = vista[vista["Stock_Actual"] <= vista["Stock_Minimo"]]
        vista["Estado"] = vista.apply(lambda r: estado_stock(r["Stock_Actual"], r["Stock_Minimo"]), axis=1)
        st.caption("💡 Edita **Precio_Actual, Stock_Actual y Stock_Minimo** directo en la tabla. "
                   "`Precio_Sugerido` se recalcula solo (Costo de Compra × 1.5). No edites el SKU.")
        edit = st.data_editor(
            vista[["SKU", "Nombre", "Categoria", "Costo", "Precio_Sugerido",
                   "Precio_Actual", "Stock_Actual", "Stock_Minimo", "Estado"]],
            use_container_width=True, num_rows="fixed",
            column_config={
                "SKU": st.column_config.TextColumn("SKU", disabled=True),
                "Categoria": st.column_config.SelectboxColumn("Categoría", options=CATEGORIAS, required=True),
                "Precio_Sugerido": st.column_config.NumberColumn("Precio Sugerido (auto)", disabled=True, format="Gs. %.0f"),
                "Precio_Actual": st.column_config.NumberColumn("Precio Actual (venta)", format="Gs. %.0f"),
                "Costo": st.column_config.NumberColumn("Costo Compra", format="Gs. %.0f"),
                "Stock_Actual": st.column_config.NumberColumn("Stock", format="%d"),
                "Stock_Minimo": st.column_config.NumberColumn("Mínimo", format="%d"),
                "Estado": st.column_config.TextColumn("Estado", disabled=True),
            },
            key="editor_inv",
        )
        col_g, col_d = st.columns([1, 1])
        if col_g.button("💾 Guardar cambios", type="primary"):
            base = inv.set_index("SKU")
            for _, r in edit.iterrows():
                sku = str(r["SKU"])
                if sku in base.index:
                    base.loc[sku, ["Nombre", "Categoria", "Costo", "Precio_Actual",
                                   "Stock_Actual", "Stock_Minimo"]] = [
                        r["Nombre"], r["Categoria"], a_entero(r["Costo"] or 0),
                        a_entero(r["Precio_Actual"] or 0), int(r["Stock_Actual"] or 0),
                        int(r["Stock_Minimo"] or 0)]
            guardar_inventario(base.reset_index())
            st.success("✅ Inventario actualizado (sugeridos recalculados desde Costo).")
            st.rerun()
        with col_d.expander("🗑️ Eliminar producto"):
            if len(inv):
                sku_del = st.selectbox("SKU a eliminar", inv["SKU"].tolist())
                if st.button("Eliminar definitivamente"):
                    guardar_inventario(inv[inv["SKU"] != sku_del])
                    st.success(f"Eliminado `{sku_del}`.")
                    st.rerun()

        st.subheader("Vista con alertas de color")
        st.dataframe(resaltar_stock(vista), use_container_width=True)
        di1, di2 = st.columns(2)
        di1.download_button("⬇️ Descargar inventario (.xlsx)",
                            df_a_excel_bytes({"Inventario": inv[COL_INV]}),
                            file_name=f"inventario-MM-{date.today()}.xlsx",
                            key="dl_inv_xlsx")
        di2.download_button("⬇️ Descargar inventario (.csv con ;)",
                            csv_puntoycoma_bytes(inv[COL_INV]),
                            file_name=f"inventario-MM-{date.today()}.csv",
                            mime="text/csv",
                            key="dl_inv_csv")

# ----------------------------------------------------------- VENTAS Y CAJA
elif pg == "🧾 Ventas y Caja":
    st.header("🧾 Ventas y Caja (Gs.)")
    if not len(inv):
        st.warning("Primero carga productos en 📦 Inventario.")
        st.stop()
    _tabs_ventas = (["🛒 Nueva venta", "📜 Historial y totales", "🗑️ Historial y Borrado"]
                    if puede_editar_inventario() else ["🛒 Nueva venta", "📜 Historial y totales"])
    if len(_tabs_ventas) == 3:
        t1, t2, t3v = st.tabs(_tabs_ventas)
    else:
        t1, t2 = st.tabs(_tabs_ventas)
        t3v = None

    with t1:
        # --- Ticket recién cobrado: vista de impresión + disparo automático ---
        if st.session_state.ultimo_ticket:
            t = st.session_state.ultimo_ticket
            auto = bool(st.session_state.auto_print_ticket)
            st.session_state.auto_print_ticket = False
            st.success(f"✅ Venta {t['id']} registrada por {fmt_gs(t['total'])} "
                       f"({t['metodo']} · {t['caja']}). Stock descontado.")
            # Estado del respaldo post-cobro a Drive (no bloqueante: el cobro
            # ya está a salvo en local; esto solo informa la subida en fondo).
            if drive_sync.modo_nube():
                _est = drive_sync.estado_ultimo_respaldo()
                try:
                    _ts_ticket = datetime.strptime(
                        str(t.get("fecha", "")), "%Y-%m-%d %H:%M").timestamp()
                except (ValueError, TypeError):
                    _ts_ticket = 0.0
                if _est and float(_est.get("ts", 0)) >= _ts_ticket - 5:
                    if _est.get("ok"):
                        st.caption(f"☁️ Respaldo en Drive completado: {_est.get('detalle', '')}")
                    else:
                        st.warning("⚠️ No se pudo respaldar en Drive en 2do plano "
                                   f"({_est.get('detalle', 'sin detalle')}). Usa el botón "
                                   "«🔄 Sincronizar con Drive» como respaldo.")
                elif st.session_state.get("_respaldo_lanzado"):
                    st.caption("☁️ Respaldando ventas y stock en Drive en 2do plano…")
            components.html(ticket_a_html_doc(t, auto_print=auto),
                            height=min(620, 300 + 44 * max(len(t["items"]), 1)),
                            scrolling=True)
            bi1, bi2 = st.columns(2)
            if bi1.button("🖨️ Reimprimir ticket", key="ticket_reprint"):
                st.session_state.auto_print_ticket = True
                st.rerun()
            if bi2.button("✖️ Cerrar ticket / nueva venta", key="ticket_close"):
                st.session_state.ultimo_ticket = None
                st.rerun()
            st.divider()
        # Función auxiliar para agregar al carrito con consolidación de ítems
        def agregar_al_carrito(sku_item: str, cant_item: int) -> tuple[bool, str]:
            sku_clean = str(sku_item).strip()
            if not sku_clean:
                return False, "Ingresa un SKU o código de barras válido."
            coincidencias = inv[inv["SKU"].astype(str) == sku_clean]
            if len(coincidencias) == 0:
                return False, f"El producto con código/SKU `{sku_clean}` no existe en inventario."
            fila_prod = coincidencias.iloc[0]
            stock_disp = int(fila_prod["Stock_Actual"])
            cant_actual = sum(it["Cantidad"] for it in st.session_state.carrito if str(it["SKU"]) == sku_clean)
            if stock_disp < 9000 and (cant_actual + cant_item) > stock_disp:
                return False, (f"Stock insuficiente para {fila_prod['Nombre']}: disponible {stock_disp}, "
                               f"en carrito {cant_actual}, intentas agregar {cant_item}.")
            # Consolidar si ya está en el ticket
            encontrado = False
            for it in st.session_state.carrito:
                if str(it["SKU"]) == sku_clean:
                    it["Cantidad"] += cant_item
                    encontrado = True
                    break
            if not encontrado:
                st.session_state.carrito.append({
                    "SKU": sku_clean,
                    "Nombre": str(fila_prod["Nombre"]),
                    "Cantidad": int(cant_item),
                    "Precio": a_entero(fila_prod["Precio_Actual"])
                })
            return True, f"Agregado: {fila_prod['Nombre']} x{cant_item}"

        st.subheader("Armar ticket")

        # Entrada rápida por lector de código de barras o SKU
        with st.form("pos_barcode_form", clear_on_submit=True):
            sb1, sb2, sb3 = st.columns([3, 1, 1])
            sku_leido = sb1.text_input("⚡ Lector de código de barras / SKU rápido",
                                       placeholder="Escanea con el lector o escribe el SKU y presiona Enter",
                                       key="pos_sku_input")
            cant_leida = sb2.number_input("Cant.", min_value=1, step=1, value=1, format="%d", key="pos_cant_input")
            if sb3.form_submit_button("⚡ Agregar rápido", type="primary", use_container_width=True):
                if sku_leido.strip():
                    ok_b, msg_b = agregar_al_carrito(sku_leido.strip(), int(cant_leida))
                    if ok_b:
                        st.success(f"✅ {msg_b}")
                        st.rerun()
                    else:
                        st.error(f"❌ {msg_b}")

        # Selección manual desde catálogo
        c1, c2, c3 = st.columns([2, 1, 1])
        opciones = {f"{r.SKU} — {r.Nombre} ({fmt_gs(r.Precio_Actual)}, stock {int(r.Stock_Actual)})": r.SKU
                    for r in inv.itertuples()}
        sel = c1.selectbox("O buscar en la lista (SKU — nombre)", list(opciones.keys()))
        cant = c2.number_input("Cantidad", min_value=1, step=1, value=1, format="%d", key="pos_manual_cant")
        sku_sel = opciones[sel]
        fila = inv[inv["SKU"] == sku_sel].iloc[0]
        c3.metric("Subtotal", fmt_gs(a_entero(fila["Precio_Actual"]) * int(cant)))
        if st.button("➕ Agregar al ticket", key="pos_btn_add_manual"):
            ok_m, msg_m = agregar_al_carrito(sku_sel, int(cant))
            if ok_m:
                st.success(f"✅ {msg_m}")
                st.rerun()
            else:
                st.error(f"❌ {msg_m}")
        if st.session_state.carrito:
            st.divider()
            car = pd.DataFrame(st.session_state.carrito)
            car["Subtotal"] = car["Cantidad"] * car["Precio"]
            mostrar = car.copy()
            mostrar["Precio"] = mostrar["Precio"].map(fmt_gs)
            mostrar["Subtotal"] = mostrar["Subtotal"].map(fmt_gs)
            st.dataframe(mostrar, use_container_width=True)
            st.caption("¿Te equivocaste en un producto? Quítalo sin vaciar todo el ticket:")
            for i, _it in enumerate(list(st.session_state.carrito)):
                if st.button(f"🗑️ Quitar «{_it['Nombre']}» x{int(_it['Cantidad'])} "
                             f"({fmt_gs(int(_it['Cantidad']) * int(_it['Precio']))})",
                             key=f"ticket_del_{i}"):
                    st.session_state.carrito.pop(i)
                    st.rerun()
            total = int(car["Subtotal"].sum())
            st.markdown(f"### Total: {fmt_gs(total)}")
            cc1, cc2, cc3 = st.columns(3)
            metodo = cc1.selectbox("Método de pago", METODOS_PAGO)
            caja = cc2.selectbox("Caja", CAJAS)
            notas = cc3.text_input("Notas (opcional)")
            ccc1, ccc2 = st.columns(2)
            if ccc1.button("✅ Cobrar e imprimir ticket", type="primary"):
                ticket = f"V-{datetime.now().strftime('%Y%m%d-%H%M%S')}"
                items_ticket = [{"SKU": it["SKU"], "Nombre": it["Nombre"],
                                 "Cantidad": int(it["Cantidad"]), "Precio": int(it["Precio"])}
                                for it in st.session_state.carrito]
                nuevas = [[ticket, ahora_str(), it["SKU"], it["Nombre"],
                           str(inv.loc[inv['SKU'] == it['SKU'], 'Categoria'].iloc[0]
                               if (inv['SKU'] == it['SKU']).any() else "Otros"),
                           int(it["Cantidad"]), int(it["Precio"]), int(it["Cantidad"] * it["Precio"]),
                           metodo, caja, notas,
                           st.session_state.auth["Documento"], st.session_state.auth["Nombre"]]
                          for it in st.session_state.carrito]
                guardar_ventas(pd.concat([ventas, pd.DataFrame(nuevas, columns=COL_VENTAS)],
                                         ignore_index=True))
                base = inv.set_index("SKU")
                for it in st.session_state.carrito:
                    if it["SKU"] in base.index and int(base.loc[it["SKU"], "Stock_Actual"]) < 9000:
                        base.loc[it["SKU"], "Stock_Actual"] = max(
                            0, int(base.loc[it["SKU"], "Stock_Actual"]) - int(it["Cantidad"]))
                guardar_inventario(base.reset_index())
                # Respaldo inmediato en la nube (solo ventas+stock, en 2do plano
                # para no bloquear). El botón manual queda como respaldo.
                try:
                    st.session_state["_respaldo_lanzado"] = bool(
                        drive_sync.respaldo_venta_fondo())
                except Exception:
                    st.session_state["_respaldo_lanzado"] = False
                st.session_state.carrito = []
                # Guardar ticket y disparar la impresión del navegador una vez
                st.session_state.ultimo_ticket = {
                    "id": ticket, "fecha": ahora_str(), "items": items_ticket,
                    "total": total, "metodo": metodo, "caja": caja, "notas": notas,
                    "vendedor": f"{st.session_state.auth['Nombre']}"}
                st.session_state.auto_print_ticket = True
                st.rerun()
            if ccc2.button("🧹 Vaciar ticket"):
                st.session_state.carrito = []
                st.rerun()
        else:
            st.info("El ticket está vacío. Agrega productos arriba.")

    with t2:
        st.subheader("Historial")
        v = ventas.copy()
        if not len(v):
            st.info("Sin ventas todavía.")
        else:
            v["Fecha"] = pd.to_datetime(v["Fecha"], errors="coerce")
            c1, c2, c3 = st.columns(3)
            d1 = c1.date_input("Desde", v["Fecha"].min().date())
            d2 = c2.date_input("Hasta", v["Fecha"].max().date())
            fm = c3.selectbox("Método", ["Todos"] + sorted(v["Metodo_Pago"].dropna().unique().tolist()))
            f = v[(v["Fecha"].dt.date >= d1) & (v["Fecha"].dt.date <= d2)]
            if fm != "Todos":
                f = f[f["Metodo_Pago"] == fm]
            ver = f.copy()
            ver["Precio_Unitario"] = ver["Precio_Unitario"].map(fmt_gs)
            ver["Total"] = ver["Total"].map(fmt_gs)
            if "Vendedor_Nombre" in ver.columns:
                ver["Vendedor"] = (ver["Vendedor_Nombre"].fillna("").astype(str) + " ("
                                   + ver["Vendedor_Doc"].fillna("").astype(str) + ")")
                ver.loc[ver["Vendedor"] == " ()", "Vendedor"] = "—"
            st.dataframe(ver.sort_values("Fecha", ascending=False), use_container_width=True)
            m1, m2, m3 = st.columns(3)
            m1.metric("Total período", fmt_gs(f["Total"].sum()))
            m2.metric("Tickets", f"{f['ID_Venta'].nunique():,}".replace(",", "."))
            m3.metric("Ticket promedio",
                      fmt_gs(f.groupby("ID_Venta")["Total"].sum().mean()) if len(f) else fmt_gs(0))
            st.download_button("⬇️ Descargar reporte de ventas (.xlsx)",
                               df_a_excel_bytes({"Ventas": f[COL_VENTAS]}),
                               file_name=f"ventas-MM-{d1}_{d2}.xlsx")

    if t3v is not None:
        with t3v:
            # Sub-pestaña restringida: solo Admin / SuperUsuario (Cajero ni la ve)
            if not puede_editar_inventario():
                st.error("🔒 Acceso restringido: solo Admin o SuperUsuario pueden borrar ventas.")
                st.stop()
            st.subheader("🗑️ Historial y Borrado de Ventas")
            st.caption("Cada ticket eliminado descuenta sus líneas de `ventas.xlsx`, "
                       "devuelve las cantidades al stock y sincroniza con Drive de inmediato.")
            vb = ventas.copy()
            if not len(vb):
                st.info("Sin ventas registradas.")
            else:
                vb["Fecha"] = pd.to_datetime(vb["Fecha"], errors="coerce")
                resumen = vb.groupby("ID_Venta").agg(
                    Fecha=("Fecha", "min"), Items=("SKU", "count"), Total=("Total", "sum"),
                    Metodo=("Metodo_Pago", "first"),
                    Vendedor=("Vendedor_Nombre", lambda s: s.fillna("").iloc[0] if len(s) else "")).reset_index()
                resumen = resumen.sort_values("Fecha", ascending=False)
                ver_r = resumen.copy()
                ver_r["Total"] = ver_r["Total"].map(fmt_gs)
                st.dataframe(ver_r, use_container_width=True)
                sel_t = st.selectbox(
                    "Ticket a revisar / borrar",
                    resumen.apply(lambda r: f"{r['ID_Venta']} — {fmt_gs(r['Total'])} ({int(r['Items'])} ítems)",
                                  axis=1).tolist(),
                    key="venta_del_sel")
                id_t = sel_t.split(" — ")[0]
                lineas = vb[vb["ID_Venta"] == id_t].copy()
                det = lineas[["SKU", "Producto", "Cantidad", "Precio_Unitario", "Total"]].copy()
                det["Precio_Unitario"] = det["Precio_Unitario"].map(fmt_gs)
                det["Total"] = det["Total"].map(fmt_gs)
                st.dataframe(det, use_container_width=True)
                f0 = lineas.iloc[0]
                _vn, _vd = f0.get("Vendedor_Nombre"), f0.get("Vendedor_Doc")
                _vend_txt = ("—" if pd.isna(_vn) or str(_vn).strip() == ""
                             else f"{_vn} ({'—' if pd.isna(_vd) else _vd})")
                st.caption(f"Fecha: {f0['Fecha']} · Método: {f0['Metodo_Pago']} · Caja: {f0['Caja']} · "
                           f"Vendedor: {_vend_txt} · Total: {fmt_gs(lineas['Total'].sum())}")
                conf_b = st.checkbox(f"Confirmo eliminar el ticket {id_t} y devolver su stock",
                                     key="venta_del_conf")
                if st.button("🗑️ Eliminar venta definitivamente", type="primary", key="venta_del_go"):
                    if not conf_b:
                        st.warning("Marca la confirmación primero.")
                    else:
                        # 1. Devolver cantidades al stock (misma regla que la venta: servicios no se tocan)
                        base = inv.set_index("SKU")
                        devueltos = 0
                        for _, r in lineas.iterrows():
                            sku = str(r["SKU"])
                            if sku in base.index and int(base.loc[sku, "Stock_Actual"]) < 9000:
                                base.loc[sku, "Stock_Actual"] = (
                                    int(base.loc[sku, "Stock_Actual"]) + int(r["Cantidad"]))
                                devueltos += int(r["Cantidad"])
                        guardar_inventario(base.reset_index())
                        # 2. Borrar las líneas (sincroniza ventas.xlsx + Drive vía hook)
                        guardar_ventas(ventas[ventas["ID_Venta"] != id_t].reset_index(drop=True))
                        st.success(f"✅ Ticket {id_t} eliminado ({fmt_gs(lineas['Total'].sum())}). "
                                   f"Stock devuelto: {devueltos} unidad(es). Drive sincronizado.")
                        st.rerun()

# ---------------------------------------------------------------- DEUDORES
elif pg == "🤝 Deudores":
    st.header("🤝 Control de deudores (Gs.)")
    t1, t2, t3 = st.tabs(["📋 Cuentas", "➕ Nuevo / Abono", "✏️ Editar deuda"])
    with t2:
        st.subheader("🧾 Nueva deuda por productos")
        c1, c2 = st.columns(2)
        cli = c1.text_input("Cliente *", key="deud_cli")
        tel = c2.text_input("Teléfono", key="deud_tel")
        c3, c4 = st.columns(2)
        fp = c3.date_input("Fecha préstamo", date.today(), key="deud_fp")
        fv = c4.date_input("Vencimiento", date.today(), key="deud_fv")

        st.markdown("**Agregar productos a la deuda**")
        if len(inv):
            opc_d = {f"{r.SKU} — {r.Nombre} ({fmt_gs(r.Precio_Actual)})": (str(r.SKU), str(r.Nombre), a_entero(r.Precio_Actual))
                     for r in inv.itertuples()}
            opc_d["✍️ Ítem manual (no está en inventario)"] = ("MANUAL", "", 0)
        else:
            opc_d = {"✍️ Ítem manual (no está en inventario)": ("MANUAL", "", 0)}
        d1, d2, d3, d4 = st.columns([3, 1, 1, 1])
        sel_d = d1.selectbox("Producto", list(opc_d.keys()), key="deud_prod_sel")
        _sku_d, _nom_d, _pre_d = opc_d[sel_d]
        nom_man = d1.text_input("Nombre del ítem manual", key="deud_man") if _sku_d == "MANUAL" else ""
        cant_d = d2.number_input("Cant.", min_value=1, step=1, value=1, format="%d", key="deud_cant")
        prec_d = d3.number_input("Precio (Gs.)", min_value=0, step=500, value=int(_pre_d), format="%d", key="deud_prec")
        d4.markdown("<br>", unsafe_allow_html=True)
        if d4.button("➕ Agregar", type="primary", use_container_width=True, key="deud_add"):
            nombre_item = nom_man.strip() if _sku_d == "MANUAL" else _nom_d
            if not nombre_item:
                st.error("Escribe el nombre del ítem manual.")
            elif a_entero(prec_d) <= 0:
                st.error("El precio debe ser mayor a 0.")
            else:
                st.session_state.items_deuda.append(
                    {"Producto": nombre_item, "Cantidad": int(cant_d),
                     "Precio": a_entero(prec_d), "Subtotal": int(cant_d) * a_entero(prec_d)})
                st.rerun()
        if st.session_state.items_deuda:
            it = pd.DataFrame(st.session_state.items_deuda)
            st.dataframe(it.assign(Precio=it["Precio"].map(fmt_gs), Subtotal=it["Subtotal"].map(fmt_gs)),
                         use_container_width=True)
            for i, row in enumerate(st.session_state.items_deuda):
                if st.button(f"❌ Quitar «{row['Producto']}» x{row['Cantidad']}", key=f"deud_del_{i}"):
                    st.session_state.items_deuda.pop(i)
                    st.rerun()
            total_d = int(sum(x["Subtotal"] for x in st.session_state.items_deuda))
            st.markdown(f"### Total deuda: {fmt_gs(total_d)}")
            nota_d = st.text_input("Nota adicional (opcional)", key="deud_nota")
            if st.button("💾 Registrar deuda", type="primary", key="deud_save"):
                if not cli.strip():
                    st.error("El Cliente es obligatorio.")
                else:
                    concepto = "; ".join(f"{x['Producto']} x{x['Cantidad']}" for x in st.session_state.items_deuda)
                    if nota_d.strip():
                        concepto += f" | Nota: {nota_d.strip()}"
                    nid = f"D-{(len(deud) + 1):04d}"
                    nuevo = pd.DataFrame([[nid, cli.strip(), tel, concepto, total_d, 0,
                                           total_d, str(fp), str(fv), "Pendiente"]], columns=COL_DEUD)
                    guardar_deudores(pd.concat([deud, nuevo], ignore_index=True))
                    st.session_state.items_deuda = []
                    st.success(f"✅ Deuda {nid} registrada por {fmt_gs(total_d)} ({len(nuevo)} cuenta).")
                    st.rerun()
            if st.button("🧹 Vaciar productos", key="deud_clear"):
                st.session_state.items_deuda = []
                st.rerun()
        else:
            st.info("Agrega al menos un producto con ➕ para formar la deuda.")

        st.divider()
        st.subheader("🔧 Agregar / quitar productos de una deuda existente")
        if len(deud):
            pendientes = deud[deud["Saldo"] > 0]
            if len(pendientes):
                sel = st.selectbox("Cuenta", pendientes.apply(
                    lambda r: f"{r['ID']} — {r['Cliente']} (saldo {fmt_gs(r['Saldo'])})", axis=1).tolist(),
                    key="deud_exist_sel")
                _id = sel.split(" — ")[0]
                fila_d = deud[deud["ID"] == _id].iloc[0]
                st.caption(f"Concepto actual: {fila_d['Concepto']} · Total: {fmt_gs(fila_d['Monto_Total'])} · "
                           f"Pagado: {fmt_gs(fila_d['Monto_Pagado'])} · Saldo: {fmt_gs(fila_d['Saldo'])}")
                e1, e2, e3 = st.columns([3, 1, 1])
                sel_e = e1.selectbox("Producto a sumar", list(opc_d.keys()), key="deud_exist_prod")
                _s2, _n2, _p2 = opc_d[sel_e]
                nom_e = ""
                if _s2 == "MANUAL":
                    nom_e = st.text_input("Nombre del ítem manual", key="deud_exist_man")
                cant_e = e2.number_input("Cant.", min_value=1, step=1, value=1, format="%d", key="deud_exist_cant")
                prec_e = e3.number_input("Precio (Gs.)", min_value=0, step=500, value=int(_p2), format="%d",
                                         key="deud_exist_prec")
                if st.button("➕ Sumar producto a la deuda", key="deud_exist_add"):
                    nombre_e = _n2 if _s2 != "MANUAL" else nom_e.strip()
                    if not nombre_e:
                        st.error("Escribe el nombre del ítem manual.")
                    elif a_entero(prec_e) <= 0:
                        st.error("El precio debe ser mayor a 0.")
                    else:
                        d = deud.copy()
                        i = d.index[d["ID"] == _id][0]
                        sub = int(cant_e) * a_entero(prec_e)
                        d.loc[i, "Concepto"] = str(d.loc[i, "Concepto"]) + f"; {nombre_e} x{int(cant_e)}"
                        d.loc[i, "Monto_Total"] = a_entero(d.loc[i, "Monto_Total"]) + sub
                        d.loc[i, "Saldo"] = a_entero(d.loc[i, "Monto_Total"]) - a_entero(d.loc[i, "Monto_Pagado"])
                        d.loc[i, "Estado"] = "Parcial" if a_entero(d.loc[i, "Monto_Pagado"]) > 0 else "Pendiente"
                        guardar_deudores(d)
                        st.success(f"✅ Sumado {nombre_e} x{int(cant_e)} ({fmt_gs(sub)}) a {_id}.")
                        st.rerun()
                with st.expander("➖ Quitar producto / descontar monto"):
                    monto_q = st.number_input("Monto a descontar (Gs.)", min_value=1000, step=1000,
                                              format="%d", key="deud_q_monto")
                    motivo_q = st.text_input("Motivo (ej. devolvió 2 cuadernos)", key="deud_q_mot")
                    if st.button("➖ Aplicar descuento", key="deud_q_go"):
                        d = deud.copy()
                        i = d.index[d["ID"] == _id][0]
                        pag = a_entero(d.loc[i, "Monto_Pagado"])
                        nuevo_tot = max(pag, a_entero(d.loc[i, "Monto_Total"]) - a_entero(monto_q))
                        d.loc[i, "Monto_Total"] = nuevo_tot
                        d.loc[i, "Saldo"] = nuevo_tot - pag
                        d.loc[i, "Concepto"] = str(d.loc[i, "Concepto"]) + (
                            f" [Quita {fmt_gs(monto_q)}: {motivo_q.strip() or 'ajuste'}]")
                        d.loc[i, "Estado"] = "Pagado" if d.loc[i, "Saldo"] <= 0 else (
                            "Parcial" if pag > 0 else "Pendiente")
                        guardar_deudores(d)
                        st.success(f"✅ Descontado {fmt_gs(monto_q)} de {_id}.")
                        st.rerun()
            else:
                st.success("Sin saldos pendientes. 🎉")
        else:
            st.info("Aún no hay deudas registradas.")

        st.divider()
        st.subheader("💵 Registrar abono / pago")
        if len(deud):
            pendientes = deud[deud["Saldo"] > 0]
            if len(pendientes):
                sel = st.selectbox("Cuenta", pendientes.apply(
                    lambda r: f"{r['ID']} — {r['Cliente']} (saldo {fmt_gs(r['Saldo'])})", axis=1).tolist())
                _id = sel.split(" — ")[0]
                monto = st.number_input("Monto del abono (Gs.)", min_value=1000, step=1000, format="%d")
                if st.button("Registrar abono"):
                    d = deud.copy()
                    i = d.index[d["ID"] == _id][0]
                    d.loc[i, "Monto_Pagado"] = a_entero(d.loc[i, "Monto_Pagado"]) + a_entero(monto)
                    d.loc[i, "Saldo"] = a_entero(d.loc[i, "Monto_Total"]) - a_entero(d.loc[i, "Monto_Pagado"])
                    if d.loc[i, "Saldo"] <= 0:
                        d.loc[i, "Saldo"] = 0
                        d.loc[i, "Estado"] = "Pagado"
                    else:
                        d.loc[i, "Estado"] = "Parcial"
                    guardar_deudores(d)
                    st.success(f"✅ Abono de {fmt_gs(monto)} registrado en {_id}.")
                    st.rerun()
            else:
                st.success("Sin saldos pendientes. 🎉")
    with t3:
        st.subheader("✏️ Editar deuda ya creada")
        if not len(deud):
            st.info("No hay deudas para editar.")
        else:
            sel_ed = st.selectbox(
                "Elige la deuda",
                deud.apply(lambda r: f"{r['ID']} — {r['Cliente']} (saldo {fmt_gs(r['Saldo'])})", axis=1).tolist(),
                key="edit_deud_sel")
            _eid = sel_ed.split(" — ")[0]
            f_ed = deud[deud["ID"] == _eid].iloc[0]
            dc1, dc2 = st.columns(2)
            ed_cli = dc1.text_input("Cliente", value=str(f_ed["Cliente"]), key=f"ed_cli_{_eid}")
            ed_tel = dc2.text_input("Teléfono", value=str(f_ed["Telefono"]), key=f"ed_tel_{_eid}")
            ed_conc = st.text_area("Concepto / productos", value=str(f_ed["Concepto"]),
                                   key=f"ed_conc_{_eid}")
            dm1, dm2, dm3 = st.columns(3)
            ed_tot = dm1.number_input("Monto Total (Gs.)", min_value=0, step=1000,
                                      value=a_entero(f_ed["Monto_Total"]), format="%d", key=f"ed_tot_{_eid}")
            ed_pag = dm2.number_input("Monto Pagado (Gs.)", min_value=0, step=1000,
                                      value=a_entero(f_ed["Monto_Pagado"]), format="%d", key=f"ed_pag_{_eid}")
            try:
                fp_def = pd.to_datetime(f_ed["Fecha_Prestamo"]).date()
            except (ValueError, TypeError):
                fp_def = date.today()
            try:
                fv_def = pd.to_datetime(f_ed["Fecha_Vencimiento"]).date()
            except (ValueError, TypeError):
                fv_def = date.today()
            df1, df2 = st.columns(2)
            ed_fp = df1.date_input("Fecha préstamo", value=fp_def, key=f"ed_fp_{_eid}")
            ed_fv = df2.date_input("Vencimiento", value=fv_def, key=f"ed_fv_{_eid}")
            nuevo_saldo = a_entero(ed_tot) - a_entero(ed_pag)
            st.info(f"Saldo recalculado: {fmt_gs(max(nuevo_saldo, 0))} "
                    f"(Total {fmt_gs(ed_tot)} − Pagado {fmt_gs(ed_pag)})")
            if st.button("💾 Guardar edición", type="primary", key=f"ed_save_{_eid}"):
                if not ed_cli.strip():
                    st.error("El Cliente no puede quedar vacío.")
                elif a_entero(ed_pag) > a_entero(ed_tot):
                    st.error("Lo pagado no puede superar el total. Ajusta los montos o registra un abono.")
                else:
                    d = deud.copy()
                    i = d.index[d["ID"] == _eid][0]
                    d.loc[i, "Cliente"] = ed_cli.strip()
                    d.loc[i, "Telefono"] = ed_tel.strip()
                    d.loc[i, "Concepto"] = ed_conc.strip() or "Deuda"
                    d.loc[i, "Monto_Total"] = a_entero(ed_tot)
                    d.loc[i, "Monto_Pagado"] = a_entero(ed_pag)
                    d.loc[i, "Saldo"] = max(a_entero(ed_tot) - a_entero(ed_pag), 0)
                    d.loc[i, "Fecha_Prestamo"] = str(ed_fp)
                    d.loc[i, "Fecha_Vencimiento"] = str(ed_fv)
                    d.loc[i, "Estado"] = ("Pagado" if d.loc[i, "Saldo"] <= 0
                                          else ("Parcial" if a_entero(ed_pag) > 0 else "Pendiente"))
                    guardar_deudores(d)
                    st.success(f"✅ Deuda {_eid} actualizada. Saldo: {fmt_gs(d.loc[i, 'Saldo'])}")
                    st.rerun()
    with t1:
        if not len(deud):
            st.info("Sin deudores. Registra el primero en la pestaña ➕ Nuevo / Abono.")
        else:
            d = deud.copy()
            d["Fecha_Vencimiento"] = pd.to_datetime(d["Fecha_Vencimiento"], errors="coerce")
            hoy = pd.Timestamp(date.today())
            d["Vencido"] = (d["Fecha_Vencimiento"] < hoy) & (d["Saldo"] > 0)
            d.loc[d["Vencido"], "Estado"] = "Vencido"
            c1, c2, c3, c4 = st.columns(4)
            c1.metric("💰 Por cobrar", fmt_gs(d["Saldo"].sum()))
            c2.metric("🔴 Vencido", fmt_gs(d[d["Vencido"]]["Saldo"].sum()))
            c3.metric("Cuentas abiertas", f"{int((d['Saldo'] > 0).sum()):,}".replace(",", "."))
            c4.metric("Cuentas pagadas", f"{int((d['Saldo'] <= 0).sum()):,}".replace(",", "."))
            fest = st.selectbox("Filtrar", ["Todas", "Pendiente", "Parcial", "Vencido", "Pagado"])
            if fest != "Todas":
                d = d[d["Estado"] == fest]
            ver = d[COL_DEUD].copy()
            for c in ("Monto_Total", "Monto_Pagado", "Saldo"):
                ver[c] = ver[c].map(fmt_gs)
            st.dataframe(ver, use_container_width=True)
            st.download_button("⬇️ Descargar deudores (.xlsx)",
                               df_a_excel_bytes({"Deudores": d[COL_DEUD]}),
                               file_name=f"deudores-MM-{date.today()}.xlsx")
            with st.expander("🗑️ Eliminar una deuda"):
                del_id = st.selectbox("Deuda a eliminar", deud["ID"].tolist(), key="deud_del_id")
                conf = st.checkbox(f"Confirmo eliminar {del_id} definitivamente", key="deud_del_conf")
                if st.button("Eliminar definitivamente", key="deud_del_go"):
                    if not conf:
                        st.warning("Marca la confirmación primero.")
                    else:
                        guardar_deudores(deud[deud["ID"] != del_id].reset_index(drop=True))
                        st.success(f"Deuda {del_id} eliminada.")
                        st.rerun()

# ------------------------------------------------------- IMPORTAR/EXPORTAR
elif pg == "🔄 Importar / Exportar":
    st.header("🔄 Importar Excel viejo / Exportar actualizado")
    st.info("Acepta `.xlsx`, `.xlsm` y `.xls`. La app detecta columnas aunque se llamen "
            "distinto (ej. `Código`→SKU, `Costo_Compra`→Costo). Montos en Gs. sin decimales.")

    c1, c2, c3 = st.columns(3)
    with c1:
        st.subheader("📦 Importar inventario")
        f = st.file_uploader("Archivo de productos", type=["xlsx", "xlsm", "xls"], key="imp_inv")
        if f:
            crudo = pd.read_excel(f, sheet_name=0, dtype=str)
            st.write("Vista previa del archivo:", crudo.head())
            norm = normalizar_df_inventario(crudo)
            st.write("Cómo quedará normalizado:", norm.head())
            modo = st.radio("Modo", ["Agregar/actualizar por SKU", "Reemplazar todo"], key="m_inv")
            if st.button("Importar inventario", type="primary"):
                if modo == "Reemplazar todo":
                    guardar_inventario(norm)
                else:
                    base = inv.set_index("SKU") if len(inv) else pd.DataFrame(columns=COL_INV).set_index("SKU")
                    for _, r in norm.set_index("SKU").iterrows():
                        base.loc[r.name] = r
                    guardar_inventario(base.reset_index())
                st.success(f"✅ {len(norm)} productos importados.")
                st.rerun()
    with c2:
        st.subheader("🧾 Importar ventas históricas")
        f = st.file_uploader("Ej. Ventas-Enero.xlsx", type=["xlsx", "xlsm", "xls"], key="imp_ven")
        if f:
            crudo = pd.read_excel(f, sheet_name=0, dtype=str)
            st.write("Vista previa:", crudo.head())
            norm = normalizar_df_ventas(crudo, inv)
            st.write("Normalizado:", norm.head())
            descontar = st.checkbox("Descontar stock al importar", value=False)
            if st.button("Importar ventas", type="primary"):
                guardar_ventas(pd.concat([ventas, norm], ignore_index=True))
                if descontar:
                    base = inv.set_index("SKU")
                    for _, r in norm.iterrows():
                        if r["SKU"] in base.index:
                            base.loc[r["SKU"], "Stock_Actual"] = max(
                                0, int(base.loc[r["SKU"], "Stock_Actual"]) - int(r["Cantidad"]))
                    guardar_inventario(base.reset_index())
                st.success(f"✅ {len(norm)} líneas importadas ({fmt_gs(norm['Total'].sum())}).")
                st.rerun()
    with c3:
        st.subheader("🤝 Importar deudores")
        f = st.file_uploader("Ej. Deudores.xlsm", type=["xlsx", "xlsm", "xls"], key="imp_deu")
        if f:
            crudo = pd.read_excel(f, sheet_name=0, dtype=str)
            st.write("Vista previa:", crudo.head())
            norm = normalizar_df_deudores(crudo)
            st.write("Normalizado:", norm.head())
            if st.button("Importar deudores", type="primary"):
                guardar_deudores(pd.concat([deud, norm], ignore_index=True))
                st.success(f"✅ {len(norm)} deudores importados.")
                st.rerun()

    st.divider()
    st.subheader("📤 Exportar todo (respaldo)")
    st.download_button("⬇️ Descargar RESPALDO COMPLETO (.xlsx: Inventario+Ventas+Deudores)",
                       df_a_excel_bytes({"Inventario": inv[COL_INV], "Ventas": ventas[COL_VENTAS],
                                         "Deudores": deud[COL_DEUD]}),
                       file_name=f"respaldo-libreria-MM-{date.today()}.xlsx")

# ---------------------------------------------------------------- USUARIOS
elif pg == "👥 Usuarios":
    if not es_gestor_usuarios():
        st.error("🔒 No tienes permiso para gestionar usuarios.")
        st.stop()
    st.header("👥 Gestión de usuarios")
    st.caption("Contraseñas guardadas con hash SHA-256. Solo SuperUsuario puede crear "
               "otro SuperUsuario o eliminar usuarios. Un SuperUsuario solo se elimina "
               "si existe otro SuperUsuario activo.")
    u = cargar_usuarios()
    ut1, ut2, ut3 = st.tabs(["📋 Usuarios", "➕ Nuevo usuario", "✏️ Editar usuario"])
    with ut2:
        roles_permitidos = ROLES if rol_actual() == "SuperUsuario" else ["Admin", "Cajero"]
        with st.form("nuevo_usu", clear_on_submit=True):
            uc1, uc2 = st.columns(2)
            doc_n = uc1.text_input("Documento (Cédula) *")
            nom_n = uc2.text_input("Nombre y apellido *")
            up1, up2, up3 = st.columns(3)
            pass1 = up1.text_input("Contraseña *", type="password")
            pass2 = up2.text_input("Repetir contraseña *", type="password")
            rol_n = up3.selectbox("Rol", roles_permitidos)
            if st.form_submit_button("💾 Crear usuario", type="primary"):
                docs = set(u["Documento"].astype(str)) if len(u) else set()
                if not doc_n.strip() or not nom_n.strip():
                    st.error("Documento y Nombre son obligatorios.")
                elif doc_n.strip() in docs:
                    st.error(f"El documento `{doc_n.strip()}` ya existe.")
                elif not puede_crear_rol(rol_n):
                    st.error("⛔ Un Admin no puede crear usuarios SuperUsuario (reservado a SuperUsuarios).")
                elif len(pass1) < 4:
                    st.error("La contraseña debe tener al menos 4 caracteres.")
                elif pass1 != pass2:
                    st.error("Las contraseñas no coinciden.")
                else:
                    nuevo = pd.DataFrame([[doc_n.strip(), nom_n.strip(), hash_password(pass1),
                                           rol_n, True, ahora_str()]], columns=COL_USU)
                    guardar_usuarios(pd.concat([u, nuevo], ignore_index=True))
                    st.success(f"✅ Usuario `{nom_n.strip()}` ({rol_n}) creado.")
                    st.rerun()
    with ut3:
        st.subheader("✏️ Editar usuario existente")
        st.caption("Modifica nombre, rol y/o contraseña. Dejar la contraseña en blanco la conserva. "
                   "No puedes cambiar tu propio rol.")
        if not len(u):
            st.info("No hay usuarios para editar.")
        else:
            sel_e = st.selectbox("Usuario a editar", u.apply(
                lambda r: f"{r['Documento']} — {r['Nombre']} ({r['Rol']})", axis=1).tolist(),
                key="usu_edit_sel")
            doc_e = sel_e.split(" — ")[0]
            f_e = u[u["Documento"].astype(str) == doc_e].iloc[0]
            es_mismo_e = doc_e == st.session_state.auth["Documento"]
            roles_ed = ROLES if rol_actual() == "SuperUsuario" else ["Admin", "Cajero"]
            ee1, ee2 = st.columns(2)
            nom_e = ee1.text_input("Nombre y apellido", value=str(f_e["Nombre"]), key=f"ue_nom_{doc_e}")
            rol_actual_e = str(f_e["Rol"])
            if es_mismo_e:
                ee2.info(f"Tu rol actual: **{rol_actual_e}** (no puedes cambiar tu propio rol)")
                rol_e = rol_actual_e
            else:
                rol_e = ee2.selectbox("Rol", roles_ed,
                                      index=roles_ed.index(rol_actual_e) if rol_actual_e in roles_ed else 0,
                                      key=f"ue_rol_{doc_e}")
            ep1, ep2 = st.columns(2)
            npw1 = ep1.text_input("Nueva contraseña (vacío = conservar)", type="password", key=f"ue_p1_{doc_e}")
            npw2 = ep2.text_input("Repetir nueva contraseña", type="password", key=f"ue_p2_{doc_e}")
            if st.button("💾 Guardar edición", type="primary", key=f"ue_save_{doc_e}"):
                if not nom_e.strip():
                    st.error("El Nombre no puede quedar vacío.")
                elif not puede_crear_rol(rol_e):
                    st.error("⛔ Un Admin no puede asignar el rol SuperUsuario (reservado a SuperUsuarios).")
                elif rol_actual_e == "SuperUsuario" and rol_e != "SuperUsuario" and not es_mismo_e \
                        and len(super_activos(u, excluir_doc=doc_e)) == 0:
                    st.error("⛔ No puedes quitarle el rol SuperUsuario: es el último SuperUsuario activo.")
                elif npw1 or npw2:
                    if len(npw1) < 4:
                        st.error("La nueva contraseña debe tener al menos 4 caracteres.")
                    elif npw1 != npw2:
                        st.error("Las nuevas contraseñas no coinciden.")
                    else:
                        d = u.copy()
                        i = d.index[d["Documento"].astype(str) == doc_e][0]
                        d.loc[i, "Nombre"] = nom_e.strip()
                        d.loc[i, "Rol"] = rol_e
                        d.loc[i, "Password_Hash"] = hash_password(npw1)
                        guardar_usuarios(d)
                        st.success(f"✅ Usuario {doc_e} actualizado (incluye nueva contraseña).")
                        st.rerun()
                else:
                    d = u.copy()
                    i = d.index[d["Documento"].astype(str) == doc_e][0]
                    d.loc[i, "Nombre"] = nom_e.strip()
                    d.loc[i, "Rol"] = rol_e
                    guardar_usuarios(d)
                    st.success(f"✅ Usuario {doc_e} actualizado.")
                    st.rerun()
    with ut1:
        if not len(u):
            st.info("Sin usuarios (además del SuperUsuario inicial).")
        else:
            ver = u[["Documento", "Nombre", "Rol", "Fecha_Creacion"]].copy()
            ver["Activo"] = u["Activo"].map(lambda x: "✅ Sí" if x else "⛔ No")
            st.dataframe(ver, use_container_width=True)
            st.divider()
            st.subheader("🔧 Acciones sobre un usuario")
            sel_u = st.selectbox("Usuario", u.apply(
                lambda r: f"{r['Documento']} — {r['Nombre']} ({r['Rol']})", axis=1).tolist(),
                key="usu_sel")
            doc_sel = sel_u.split(" — ")[0]
            es_mismo = doc_sel == st.session_state.auth["Documento"]
            ua1, ua2, ua3 = st.columns(3)
            fila_u = u[u["Documento"].astype(str) == doc_sel].iloc[0]
            if ua1.button("🔄 Activar/Desactivar", key="usu_toggle"):
                if es_mismo:
                    st.error("No puedes desactivar tu propio usuario.")
                else:
                    d = u.copy()
                    i = d.index[d["Documento"].astype(str) == doc_sel][0]
                    va_a_desactivar = bool(d.loc[i, "Activo"])
                    if va_a_desactivar and str(d.loc[i, "Rol"]) == "SuperUsuario" \
                            and len(super_activos(u, excluir_doc=doc_sel)) == 0:
                        st.error("⛔ No puedes desactivar al último SuperUsuario activo del sistema.")
                    else:
                        d.loc[i, "Activo"] = not bool(d.loc[i, "Activo"])
                        guardar_usuarios(d)
                        st.success(f"✅ {doc_sel} ahora está {'activo' if d.loc[i, 'Activo'] else 'inactivo'}.")
                        st.rerun()
            with ua2.expander("🔑 Restablecer contraseña"):
                np1 = st.text_input("Nueva contraseña", type="password", key="usu_np1")
                np2 = st.text_input("Repetir", type="password", key="usu_np2")
                if st.button("Guardar contraseña", key="usu_np_go"):
                    if len(np1) < 4:
                        st.error("Mínimo 4 caracteres.")
                    elif np1 != np2:
                        st.error("No coinciden.")
                    else:
                        d = u.copy()
                        i = d.index[d["Documento"].astype(str) == doc_sel][0]
                        d.loc[i, "Password_Hash"] = hash_password(np1)
                        guardar_usuarios(d)
                        st.success(f"✅ Contraseña de {doc_sel} actualizada.")
                        st.rerun()
            with ua3.expander("🗑️ Eliminar usuario"):
                if rol_actual() != "SuperUsuario":
                    st.info("Solo SuperUsuario puede eliminar usuarios.")
                elif es_mismo:
                    st.info("No puedes eliminar tu propio usuario.")
                else:
                    fila_del = u[u["Documento"].astype(str) == doc_sel].iloc[0]
                    if str(fila_del["Rol"]) == "SuperUsuario" \
                            and len(super_activos(u, excluir_doc=doc_sel)) == 0:
                        st.error("⛔ No se puede eliminar: es el último SuperUsuario activo. "
                                 "Crea otro SuperUsuario primero.")
                    elif st.button(f"Eliminar {doc_sel} definitivamente", key="usu_del_go"):
                        guardar_usuarios(u[u["Documento"].astype(str) != doc_sel].reset_index(drop=True))
                        st.success(f"Usuario {doc_sel} eliminado.")
                        st.rerun()

# ---------------------------------------------------------------- REPORTES
else:
    st.header("📈 Reportes (Gs.)")
    v = ventas.copy()
    if not len(v):
        st.info("Sin ventas para reportar todavía.")
        st.stop()
    v["Fecha"] = pd.to_datetime(v["Fecha"], errors="coerce")
    v = v.dropna(subset=["Fecha"])
    c1, c2 = st.columns(2)
    d1 = c1.date_input("Desde", v["Fecha"].min().date())
    d2 = c2.date_input("Hasta", v["Fecha"].max().date())
    f = v[(v["Fecha"].dt.date >= d1) & (v["Fecha"].dt.date <= d2)]
    if not len(f):
        st.warning("Sin ventas en ese rango.")
        st.stop()
    k1, k2, k3, k4 = st.columns(4)
    k1.metric("Total", fmt_gs(f["Total"].sum()))
    k2.metric("Unidades", f"{int(f['Cantidad'].sum()):,}".replace(",", "."))
    k3.metric("Tickets", f"{f['ID_Venta'].nunique():,}".replace(",", "."))
    k4.metric("Ticket promedio", fmt_gs(f.groupby("ID_Venta")["Total"].sum().mean()))
    st.subheader("Evolución en el período")
    evo = f.groupby(f["Fecha"].dt.date)["Total"].sum().reset_index()
    evo.columns = ["Fecha", "Total"]
    evo["Fecha"] = pd.to_datetime(evo["Fecha"])
    st.altair_chart(graf_serie(evo, "Fecha", "Total", alto=280), use_container_width=True)
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("Por categoría")
        st.altair_chart(graf_barras(f.groupby("Categoria", as_index=False)["Total"].sum()
                                    .sort_values("Total", ascending=False),
                                    "Categoria", "Total", horizontal=True),
                        use_container_width=True)
        st.subheader("Por método de pago")
        st.altair_chart(graf_dona(f.groupby("Metodo_Pago", as_index=False)["Total"].sum(),
                                  "Metodo_Pago", "Total"),
                        use_container_width=True)
    with c2:
        st.subheader("Por caja")
        st.altair_chart(graf_barras(f.groupby("Caja", as_index=False)["Total"].sum(),
                                    "Caja", "Total"),
                        use_container_width=True)
        if "Vendedor_Nombre" in f.columns and f["Vendedor_Nombre"].fillna("").ne("").any():
            st.subheader("Por vendedor (auditoría)")
            st.altair_chart(graf_barras(f[f["Vendedor_Nombre"].fillna("") != ""]
                                        .groupby("Vendedor_Nombre", as_index=False)["Total"].sum()
                                        .sort_values("Total", ascending=False),
                                        "Vendedor_Nombre", "Total", horizontal=True),
                            use_container_width=True)
        st.subheader("Top 10 productos")
        top = f.groupby("Producto").agg(Unidades=("Cantidad", "sum"), Total=("Total", "sum")).sort_values(
            "Total", ascending=False).head(10)
        top_ver = top.copy()
        top_ver["Total"] = top_ver["Total"].map(fmt_gs)
        st.dataframe(top_ver, use_container_width=True)
    if len(inv):
        costos = dict(zip(inv["SKU"].astype(str), inv["Costo"]))
        ff = f.copy()
        ff["Costo_U"] = ff["SKU"].map(costos).fillna(0).astype(int)
        ff["Ganancia"] = (ff["Precio_Unitario"] - ff["Costo_U"]) * ff["Cantidad"]
        st.metric("💹 Ganancia bruta estimada del período", fmt_gs(ff["Ganancia"].sum()))
    st.download_button("⬇️ Descargar reporte del período (.xlsx)",
                       df_a_excel_bytes({"Resumen": f.groupby("Producto").agg(
                           Unidades=("Cantidad", "sum"), Total=("Total", "sum")).reset_index(),
                           "Detalle": f[COL_VENTAS]}),
                       file_name=f"reporte-MM-{d1}_{d2}.xlsx")

st.divider()
st.caption(f"Librería MM · datos en {DATA_DIR} · Sugerido = Costo Compra × {MARGEN} · Montos en Gs. · {ahora_str()}")
